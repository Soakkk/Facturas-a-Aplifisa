"""Recoger las facturas sueltas del Escritorio y Descargas.

Con el tiempo se amontonan PDF de facturas y Excel de Aplifisa fuera del
archivo documental. Aquí se buscan, se averigua de qué cliente, ejercicio y
tipo es cada uno y se propone llevarlos a su carpeta:

    Documentación Facturas / Nombre — NIF / Ejercicio / Gastos | Ingresos
    Documentación Facturas / Nombre — NIF / Ejercicio / Excel Aplifisa

Reglas:
  - No se mueve nada sin que la persona revise la propuesta.
  - Primero se lee el texto del PDF (gratis). Solo los escaneados sin texto
    pueden leerse con Gemini, y solo su primera página, avisando del coste.
  - El cliente se reconoce por un NIF que ya es cliente de la asesoría (aquí
    o en el directorio de la suite). Un PDF sin NIF conocido no se adivina:
    queda «sin identificar» y en su sitio.
  - Una copia idéntica de algo ya archivado no se duplica: va a _Duplicados.
  - Todo movimiento queda apuntado y se puede deshacer.
"""

from __future__ import annotations

import json
import os
import re
import shutil
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional
from uuid import uuid4

from . import archivo, clientes, identidad_archivo, suite
from .escaner import sanear
from .validacion import validar_nif

CARPETA_EXCEL = "Excel Aplifisa"
GASTOS, INGRESOS = "gastos", "ingresos"
MOVER, DUPLICADO, NO_TOCAR = "mover", "duplicado", "no_tocar"
PROFUNDIDAD = 2          # subcarpetas del Escritorio/Descargas que se miran
MAX_PAGINAS_TEXTO = 3

_RE_NIF = re.compile(
    r"(?<![A-Z0-9])(?:ES[\s.-]?)?([A-Z][\s.-]?\d{7}[\s.-]?[A-Z0-9]|\d{8}[\s.-]?[A-Z]"
    r"|[XYZ][\s.-]?\d{7}[\s.-]?[A-Z])(?![A-Z0-9])")
_RE_FECHA = re.compile(r"(?<!\d)(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})(?!\d)")
_RE_EXCEL = re.compile(r"^(GASTOS|INGRESOS)_(.+?)(?:_\d+)?\.xlsx$", re.IGNORECASE)


# ----------------------------------------------------------------- orígenes
def carpetas_origen() -> List[str]:
    """Escritorio y Descargas (también los de OneDrive), sin repetir."""
    casa = Path(os.path.expanduser("~"))
    posibles = [casa / "Desktop", casa / "Escritorio", casa / "Downloads",
                casa / "Descargas", casa / "OneDrive" / "Desktop",
                casa / "OneDrive" / "Escritorio"]
    vistas, salida = set(), []
    for ruta in posibles:
        if ruta.is_dir():
            real = os.path.normcase(os.path.realpath(ruta))
            if real not in vistas:
                vistas.add(real)
                salida.append(str(ruta))
    return salida


def clientes_conocidos(base: str) -> Dict[str, str]:
    """{NIF: nombre} de los clientes de la asesoría que se pueden reconocer."""
    conocidos: Dict[str, str] = {}
    try:
        for nif, ficha in identidad_archivo.leer(base).items():
            conocidos[clientes._normaliza(nif)] = ficha["nombre"]
    except (OSError, ValueError):
        pass
    for nif, ficha in clientes._leer_todo().items():
        if isinstance(ficha, dict) and ficha.get("confirmado") and ficha.get("nombre"):
            conocidos.setdefault(clientes._normaliza(nif), str(ficha["nombre"]).strip())
    for nif in suite.clientes():
        nombre = suite.nombre_de(nif)
        if nombre:
            conocidos.setdefault(nif, nombre)
    return {n: v for n, v in conocidos.items() if n and v}


# -------------------------------------------------------------- candidatos
@dataclass
class Candidato:
    ruta: str
    clase: str = "factura"           # factura | excel | listado
    nif: str = ""
    nombre: str = ""
    ejercicio: Optional[int] = None
    tipo: str = ""                   # gastos | ingresos
    como: str = ""                   # cómo se ha identificado
    dudas: List[str] = field(default_factory=list)
    opciones: List[str] = field(default_factory=list)   # NIF de clientes posibles
    sin_texto: bool = False          # escaneado: necesitaría Gemini
    huella: str = ""
    accion: str = NO_TOCAR
    destino: str = ""
    duplicado_de: str = ""

    @property
    def listo(self) -> bool:
        """Se puede mover sin preguntar nada más."""
        return bool(self.nif and self.ejercicio and self.tipo and not self.dudas)


def _nifs_del_texto(texto: str) -> List[str]:
    """NIF válidos en el orden en que aparecen (sin repetir)."""
    salida = []
    for encaje in _RE_NIF.finditer(texto.upper()):
        nif = re.sub(r"[\s.-]", "", encaje.group(1))
        if validar_nif(nif) and nif not in salida:
            salida.append(nif)
    return salida


def _ejercicio_del_texto(texto: str) -> Optional[int]:
    """El año que más se repite en las fechas del documento."""
    hoy = date.today()
    anos = []
    for d, m, a in _RE_FECHA.findall(texto):
        ano = int(a) + (2000 if len(a) == 2 else 0)
        if 1 <= int(d) <= 31 and 1 <= int(m) <= 12 and 2000 <= ano <= hoy.year:
            anos.append(ano)
    return Counter(anos).most_common(1)[0][0] if anos else None


def _texto_pdf(ruta: str) -> tuple[str, int]:
    import fitz
    with fitz.open(ruta) as doc:
        paginas = doc.page_count
        texto = "\n".join(doc[i].get_text() for i in range(min(paginas, MAX_PAGINAS_TEXTO)))
    return texto, paginas


def identificar_por_datos(c: Candidato, nifs: List[str], ejercicio: Optional[int],
                          conocidos: Dict[str, str], como: str,
                          emisor: str = "") -> None:
    """Rellena cliente, tipo y ejercicio a partir de los NIF encontrados.

    `nifs` va en orden de aparición. Si se sabe quién emite (lectura con IA),
    `emisor` lo dice; si no, el primer NIF del documento se toma como emisor
    (las facturas llevan arriba al que las emite) y se avisa si no hay otro.
    """
    del_cliente = [n for n in nifs if n in conocidos]
    c.opciones = del_cliente
    if not del_cliente:
        c.dudas.append("Ningún NIF del documento es de un cliente conocido")
        return
    if len(del_cliente) > 1:
        # Dos clientes de la asesoría en la misma factura (uno compra al otro):
        # no se elige a ciegas.
        c.dudas.append("Aparecen varios clientes de la asesoría: "
                       + ", ".join(f"{conocidos[n]} ({n})" for n in del_cliente))
        c.nif, c.nombre = del_cliente[0], conocidos[del_cliente[0]]
    else:
        c.nif, c.nombre = del_cliente[0], conocidos[del_cliente[0]]
    c.como = como
    if emisor:
        c.tipo = INGRESOS if emisor == c.nif else GASTOS
    elif len(nifs) >= 2:
        c.tipo = INGRESOS if nifs[0] == c.nif else GASTOS
    else:
        c.tipo = GASTOS
        c.dudas.append("Solo aparece el NIF del cliente: compruebe si es gasto o ingreso")
    c.ejercicio = ejercicio
    if not ejercicio:
        c.dudas.append("No se ha encontrado la fecha: indique el ejercicio")


def _analizar_pdf(c: Candidato, conocidos: Dict[str, str]) -> None:
    from .registro import parece_listado
    try:
        texto, _paginas = _texto_pdf(c.ruta)
    except Exception as e:  # PDF dañado o protegido
        c.dudas.append(f"No se puede abrir el PDF: {e}")
        return
    if parece_listado(c.ruta):
        c.clase = "listado"
        c.como = "Listado de Aplifisa (no es una factura)"
        return
    if len(texto.strip()) < 40:
        c.sin_texto = True
        c.dudas.append("Escaneado sin texto: léalo con Gemini o indique el cliente")
        return
    nifs = _nifs_del_texto(texto)
    identificar_por_datos(c, nifs, _ejercicio_del_texto(texto), conocidos,
                          "NIF en el texto del PDF")


def _analizar_excel(c: Candidato, conocidos: Dict[str, str]) -> None:
    """GASTOS_<cliente>.xlsx: el cliente por el nombre y el año por sus fechas."""
    encaje = _RE_EXCEL.match(os.path.basename(c.ruta))
    c.clase = "excel"
    c.tipo = GASTOS if encaje.group(1).upper() == "GASTOS" else INGRESOS
    buscado = clientes._clave_nombre(encaje.group(2))
    candidatos = [n for n, nombre in conocidos.items()
                  if clientes._clave_nombre(sanear(nombre)) == buscado
                  or clientes._clave_nombre(nombre) == buscado]
    c.opciones = candidatos
    if len(candidatos) == 1:
        c.nif, c.nombre = candidatos[0], conocidos[candidatos[0]]
        c.como = "Nombre del Excel exportado"
    else:
        c.dudas.append("No se reconoce el cliente por el nombre del Excel")
    c.ejercicio = _ejercicio_excel(c.ruta)
    if not c.ejercicio:
        c.dudas.append("El Excel no tiene fechas legibles: indique el ejercicio")


def _ejercicio_excel(ruta: str) -> Optional[int]:
    from openpyxl import load_workbook
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception:
        return None
    anos = []
    try:
        hoja = libro.active
        for fila in hoja.iter_rows(min_row=1, max_row=500, values_only=True):
            for valor in fila:
                if isinstance(valor, (datetime, date)):
                    anos.append(valor.year)
                elif isinstance(valor, str):
                    ano = _ejercicio_del_texto(valor)
                    if ano:
                        anos.append(ano)
    finally:
        libro.close()
    return Counter(anos).most_common(1)[0][0] if anos else None


def buscar(origenes: List[str], base: str,
           excluir: Optional[List[str]] = None,
           progreso: Optional[Callable[[int, int], None]] = None) -> List[Candidato]:
    """PDF y Excel de Aplifisa sueltos en los orígenes, ya analizados."""
    base_real = os.path.normcase(os.path.realpath(base))
    excluidas = {os.path.normcase(os.path.realpath(r)) for r in (excluir or [])}
    rutas: List[str] = []
    for origen in origenes:
        origen = os.path.realpath(origen)
        nivel_base = origen.rstrip(os.sep).count(os.sep)
        for raiz, carpetas, archivos in os.walk(origen):
            real = os.path.normcase(os.path.realpath(raiz))
            if real == base_real or real.startswith(base_real + os.sep):
                carpetas[:] = []
                continue
            carpetas[:] = [c for c in carpetas if not c.startswith((".", "_", "$"))
                           and raiz.count(os.sep) - nivel_base < PROFUNDIDAD]
            for nombre in archivos:
                ruta = os.path.join(raiz, nombre)
                bajo = nombre.lower()
                if (bajo.endswith(".pdf") or _RE_EXCEL.match(nombre)) and \
                        os.path.normcase(os.path.realpath(ruta)) not in excluidas:
                    rutas.append(ruta)
    conocidos = clientes_conocidos(base)
    salida = []
    for i, ruta in enumerate(sorted(set(rutas))):
        c = Candidato(ruta=ruta)
        try:
            c.huella = identidad_archivo.huella(ruta)
        except OSError as e:
            c.dudas.append(f"No se puede leer: {e}")
            salida.append(c)
            continue
        if ruta.lower().endswith(".xlsx"):
            _analizar_excel(c, conocidos)
        else:
            _analizar_pdf(c, conocidos)
        salida.append(c)
        if progreso:
            progreso(i + 1, len(rutas))
    return salida


def leer_con_gemini(candidatos: List[Candidato], base: str, api_key: str,
                    progreso: Optional[Callable[[int, int], None]] = None) -> float:
    """Identifica con Gemini (solo la 1ª página) los escaneados sin texto.

    Devuelve lo que ha costado, en euros. Una sola lectura: aquí solo hace
    falta saber de quién es, no los importes.
    """
    import fitz
    from . import costes
    from .extraccion import DOBLE_NO, Extractor
    from .pdf import CALIDAD

    conocidos = clientes_conocidos(base)
    extractor = Extractor(api_key, modo_doble=DOBLE_NO)
    pendientes = [c for c in candidatos if c.sin_texto and c.clase == "factura"]
    gasto = 0.0
    for i, c in enumerate(pendientes):
        try:
            with fitz.open(c.ruta) as doc:
                img = doc[0].get_pixmap(dpi=150).pil_tobytes(
                    format="JPEG", quality=CALIDAD)
            leido = extractor.extraer(img, c.ruta, 1)
            for modelo, entrada, salida in leido.consumos:
                gasto += costes.registrar(modelo, entrada, salida)
        except Exception as e:
            gasto += sum(costes.registrar(m, e_, s)
                         for m, e_, s in getattr(e, "consumos", []) or [])
            c.dudas = [f"Gemini no pudo leerlo: {str(e)[:80]}"]
            continue
        datos = leido.crudo
        emisor = re.sub(r"[\s.-]", "", str(datos.get("emisor_nif") or "")).upper()
        receptor = re.sub(r"[\s.-]", "", str(datos.get("receptor_nif") or "")).upper()
        from .validacion import fecha_de
        dia = fecha_de(datos.get("fecha")) if datos.get("fecha") else None
        c.sin_texto = False
        c.dudas = []
        identificar_por_datos(
            c, [n for n in (emisor, receptor) if n], dia.year if dia else None,
            conocidos, "Leído con Gemini (1ª página)", emisor=emisor)
        if not c.opciones:
            partes = [p for p in (datos.get("emisor_nombre"), datos.get("receptor_nombre")) if p]
            if partes:
                c.dudas = [f"Entre {' y '.join(partes)} no hay ningún cliente conocido"]
        if progreso:
            progreso(i + 1, len(pendientes))
    return round(gasto, 4)


def coste_estimado_gemini(candidatos: List[Candidato]) -> float:
    from . import costes
    from .extraccion import modelos_configurados
    n = sum(1 for c in candidatos if c.sin_texto and c.clase == "factura")
    return round(n * costes.coste_por_factura(150, modelos_configurados()[0]), 4)


# ------------------------------------------------------------------ plan
def _carpeta_cliente(base: str, nombre: str, nif: str) -> str:
    """Nombre de la carpeta del cliente SIN crear nada (solo se mira)."""
    try:
        indice = identidad_archivo.leer(base)
    except (OSError, ValueError):
        indice = {}
    nif = clientes._normaliza(nif)
    if nif in indice:
        return indice[nif]["carpeta"]
    return f"{sanear(nombre)} — {nif}"


def _destino_carpeta(base: str, c: Candidato) -> str:
    ejercicio = os.path.join(base, _carpeta_cliente(base, c.nombre, c.nif),
                             str(int(c.ejercicio)))
    if c.clase == "excel":
        return os.path.join(ejercicio, CARPETA_EXCEL)
    return os.path.join(ejercicio, "Ingresos" if c.tipo == INGRESOS else "Gastos")


def _nombre_libre(carpeta: str, nombre: str, ocupados: set) -> str:
    raiz, ext = os.path.splitext(nombre)
    raiz = sanear(raiz)
    ruta = os.path.join(carpeta, raiz + ext.lower())
    n = 2
    while os.path.exists(ruta) or os.path.normcase(ruta) in ocupados:
        ruta = os.path.join(carpeta, f"{raiz}_{n}{ext.lower()}")
        n += 1
    ocupados.add(os.path.normcase(ruta))
    return ruta


def planificar(candidatos: List[Candidato], base: str) -> List[Candidato]:
    """Decide qué pasa con cada candidato elegido (mover o apartar duplicado).

    Solo se planifican los que tienen cliente, ejercicio y tipo; los demás se
    quedan donde están.
    """
    ocupados: set = set()
    vistos: Dict[str, str] = {}
    for c in candidatos:
        c.destino, c.duplicado_de = "", ""
        if c.clase == "listado" or not (c.nif and c.nombre and c.ejercicio and c.tipo):
            c.accion = NO_TOCAR
            continue
        carpeta = _destino_carpeta(base, c)
        # ¿Ya está archivado (o repetido entre los propios sueltos)?
        previo = vistos.get(c.huella)
        if not previo and os.path.isdir(carpeta):
            for existente in Path(carpeta).iterdir():
                if (existente.is_file() and existente.stat().st_size == os.path.getsize(c.ruta)
                        and identidad_archivo.huella(existente) == c.huella):
                    previo = str(existente)
                    break
        if previo:
            c.accion, c.duplicado_de = DUPLICADO, previo
            carpeta_dup = os.path.join(base, identidad_archivo.DUPLICADOS,
                                       f"Recogida {date.today():%Y-%m-%d}")
            c.destino = _nombre_libre(carpeta_dup, os.path.basename(c.ruta), ocupados)
            continue
        c.accion = MOVER
        c.destino = _nombre_libre(carpeta, os.path.basename(c.ruta), ocupados)
        vistos[c.huella] = c.destino
    return candidatos


def _trasladar(origen: str, destino: str, huella: str) -> None:
    """Copia exclusiva + comprobación + borrado del original (como organizar)."""
    if identidad_archivo.huella(origen) != huella:
        raise ValueError(f"El archivo ha cambiado: {os.path.basename(origen)}.")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    creada = False
    try:
        with open(origen, "rb") as entrada, open(destino, "xb") as salida:
            creada = True
            shutil.copyfileobj(entrada, salida)
            salida.flush()
            os.fsync(salida.fileno())
        if identidad_archivo.huella(destino) != huella:
            raise OSError(f"No se pudo comprobar la copia de {os.path.basename(origen)}.")
        shutil.copystat(origen, destino)
        os.remove(origen)
    except Exception:
        if creada and os.path.exists(origen) and os.path.exists(destino):
            os.remove(destino)
        raise


def aplicar(candidatos: List[Candidato], base: str) -> dict:
    """Mueve lo planificado y lo apunta para poder deshacerlo."""
    movimientos = [{"origen": c.ruta, "destino": c.destino, "sha256": c.huella,
                    "accion": c.accion}
                   for c in candidatos if c.accion in (MOVER, DUPLICADO) and c.destino]
    # Ahora sí: se dan de alta las carpetas de los clientes que reciben algo,
    # con sus dos carpetas de Gastos e Ingresos, como al escanear.
    for c in candidatos:
        if c.accion == MOVER and c.destino:
            archivo.carpeta_tipo_cliente(c.nombre, c.ejercicio, c.tipo, base, nif=c.nif)
    registro = {"tipo": "recogida", "movimientos": movimientos, "completados": [],
                "estado": "en curso", "fecha": datetime.now().isoformat(timespec="seconds")}
    ruta = Path(base) / identidad_archivo.HISTORIAL / \
        f"{datetime.now():%Y%m%d-%H%M%S-%f}-recogida-{uuid4().hex[:6]}.json"
    identidad_archivo.guardar(ruta, registro)
    errores = []
    for m in movimientos:
        try:
            _trasladar(m["origen"], m["destino"], m["sha256"])
            registro["completados"].append(m)
        except (OSError, ValueError) as e:
            errores.append(f"{os.path.basename(m['origen'])}: {e}")
        identidad_archivo.guardar(ruta, registro)
    registro["estado"] = "completado" if not errores else "incompleto"
    registro["errores"] = errores
    identidad_archivo.guardar(ruta, registro)
    return {"registro": str(ruta),
            "movidos": sum(1 for m in registro["completados"] if m["accion"] == MOVER),
            "duplicados": sum(1 for m in registro["completados"] if m["accion"] == DUPLICADO),
            "errores": errores}


def deshacer_ultima(base: str) -> int:
    """Devuelve a su sitio lo movido en la última recogida. Cuántos volvieron."""
    carpeta = Path(base) / identidad_archivo.HISTORIAL
    for ruta in sorted(carpeta.glob("*-recogida-*.json"), reverse=True):
        registro = json.loads(ruta.read_text(encoding="utf-8"))
        if registro.get("estado") == "deshecho":
            continue
        vueltos = 0
        for m in reversed(registro.get("completados", [])):
            if os.path.exists(m["origen"]) or not os.path.exists(m["destino"]):
                continue
            _trasladar(m["destino"], m["origen"], m["sha256"])
            vueltos += 1
        registro["estado"] = "deshecho"
        identidad_archivo.guardar(ruta, registro)
        return vueltos
    return 0


def afectados(candidatos: List[Candidato]) -> List[tuple]:
    """(nombre, nif, ejercicio) de los expedientes que conviene rehacer."""
    return sorted({(c.nombre, c.nif, c.ejercicio) for c in candidatos
                   if c.accion == MOVER and c.nif and c.ejercicio})
