"""Lecturas imposibles de la IA: ni paran la cola ni se pierden facturas (1.26).

Lo que encontró el fuzzing «lecturas absurdas» (09/10/2026, semilla
20261009): un NaN, un infinito o un 10**400 donde va un importe paraban la
cola y el lote no se recuperaba; una lista donde va un nombre tiraba el bloque
entero de 25 hojas ya pagadas; cientos de líneas de IVA congelaban la ficha…
Cada prueba falla sin su arreglo. Solo datos inventados (NIF de prueba).
"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import copy
import json
import math
import time

import pytest
from PySide6.QtCore import QObject, Signal
from PySide6.QtWidgets import QApplication

from facturas_excel import clientes, procesar, sesion, suite, ventana_lectura
from facturas_excel.extraccion import _num
from facturas_excel.modelo import Factura
from facturas_excel.resumen import eur
from facturas_excel.tabla_facturas import parse_numero
from facturas_excel.validacion import ERROR, OK, REVISAR, validar

_app = QApplication.instance() or QApplication([])

NAN, INF = float("nan"), float("inf")
CLIENTE = ("ANTONIO PRUEBA EJEMPLO", "12345678Z")
PROVEEDORES = [("PROVEEDOR UNO SL", "B12345674"), ("PROVEEDOR DOS SA", "A12345674")]


@pytest.fixture(autouse=True)
def cliente_conocido():
    suite._cache.update(mtime=None, ruta=None, datos=None)
    clientes.marcar_cliente(CLIENTE[1], CLIENTE[0])


def lectura(i: int = 0, **cambios) -> dict:
    """Lo que devolvería Gemini de la factura i: completa y correcta."""
    nombre, nif = PROVEEDORES[i % len(PROVEEDORES)]
    base = round(50 + i * 37.13, 2)
    cuota = round(base * 0.21, 2)
    datos = {
        "emisor_nombre": nombre, "emisor_nif": nif,
        "receptor_nombre": CLIENTE[0], "receptor_nif": CLIENTE[1],
        "num_factura": f"F26/{100 + i:05d}", "fecha": f"{1 + i % 28:02d}/03/2026",
        "fecha_operacion": None, "estado_pagina_factura": "unica",
        "lineas_iva": [{"base": base, "tipo_iva": 21.0, "cuota_iva": cuota,
                        "pct_requiv": None, "cuota_requiv": None}],
        "base_irpf": None, "pct_irpf": None, "cuota_irpf": None, "suplidos": None,
        "es_bien_inversion": False, "total": round(base + cuota, 2),
        "sustituye_a": None, "manuscrito_en_importes": False,
        "cuenta_gasto": "629", "subclave_gxx": "G22", "cuenta_ingreso": "705",
        "subclave_ingreso": "I01", "concepto_texto": "material de oficina",
        "tipo_documento": "factura", "moneda": "EUR", "mencion_iva": "ninguna",
        "posible_no_deducible": "no", "confianza": "alta",
    }
    datos.update(cambios)
    return datos


def crudos_de(*lecturas, origen="taco.pdf"):
    return [(b"", origen, i + 1, d) for i, d in enumerate(lecturas)]


def lote(*lecturas):
    return procesar.preparar_lote(crudos_de(*lecturas), *CLIENTE)


def factura_correcta(**cambios) -> Factura:
    f = Factura(num_factura="F-1", fecha="02/03/2026", concepto="629",
                subclave="G22", nombre="PROVEEDOR UNO SL", nif="B12345674",
                base_iva=100.0, pct_iva=21.0, cuota_iva=21.0, total_impreso=121.0)
    for campo, valor in cambios.items():
        setattr(f, campo, valor)
    return f


# =====================================================================
# n.º 1 — Importes imposibles: NaN, infinito, 1e999, 10**400, un billón
# =====================================================================
IMPOSIBLES = ["NaN", "nan", "inf", "-inf", "Infinity", "1e999", NAN, INF, -INF,
              10 ** 400, -10 ** 400, 2e9, "2000000000", 1e300]


def corto(valor) -> str:
    """El nombre de un caso: un entero de 400 cifras no cabe en una línea."""
    if isinstance(valor, int) and abs(valor) > 10 ** 30:
        return f"{'-' if valor < 0 else ''}entero_de_{len(str(abs(valor)))}_cifras"
    texto = ascii(valor)
    return texto if len(texto) <= 40 else f"{texto[:30]}…{len(texto)}"


@pytest.mark.parametrize("valor", IMPOSIBLES, ids=corto)
def test_num_no_deja_pasar_un_importe_imposible(valor):
    assert _num(valor) is None


@pytest.mark.parametrize("valor", IMPOSIBLES, ids=corto)
def test_una_celda_con_un_importe_imposible_queda_vacia(valor):
    texto = valor if isinstance(valor, str) else str(valor)
    assert parse_numero(texto) is None


def test_los_importes_normales_siguen_igual():
    assert _num("1.234,56") == 1234.56
    assert _num("15,51-") == -15.51
    assert _num(1e9) == 1e9 and _num(-121) == -121.0
    assert parse_numero("1.234,56") == 1234.56
    assert parse_numero("-22,50") == -22.5


@pytest.mark.parametrize("valor", [NAN, INF, -INF, None])
def test_eur_no_lanza_con_un_importe_imposible(valor):
    assert eur(valor) == "— €"
    assert eur(-1234.5) == "-1.234,50 €"


@pytest.mark.parametrize("campo", ["base_iva", "cuota_iva", "total_impreso",
                                   "suplidos", "cuota_irpf"])
@pytest.mark.parametrize("valor", [NAN, INF, 2e9])
def test_validar_deja_en_rojo_un_importe_imposible(campo, valor):
    # Una sesión de antes de la 1.26 puede traerlo: con NaN ninguna cuenta
    # descuadraba y la fila salía en verde.
    assert validar(factura_correcta()).estado == OK
    resultado = validar(factura_correcta(**{campo: valor}))
    assert resultado.estado == ERROR
    assert any("Importe imposible" in m for m in resultado.mensajes)


def test_exportar_nunca_escribe_nan_ni_inf(tmp_path):
    from openpyxl import load_workbook

    from facturas_excel.config_columnas import leer_config
    from facturas_excel.exportar import (MODO_NUMERO, MODO_TEXTO, exportar_excel,
                                         verificar_excel)
    from facturas_excel.rutas import ruta_config

    facturas = [factura_correcta(base_iva=NAN, cuota_iva=INF, total_impreso=-INF),
                factura_correcta()]
    config = leer_config(ruta_config("gastos.xml"))
    for modo in (MODO_TEXTO, MODO_NUMERO):
        ruta = str(tmp_path / f"gastos_{modo}.xlsx")
        exportar_excel(facturas, config, ruta, modo)
        if modo == MODO_TEXTO:          # el que usa el programa
            assert verificar_excel(facturas, config, ruta, modo) == []
        libro = load_workbook(ruta)
        celdas = [str(x).lower() for fila in libro.active.iter_rows(values_only=True)
                  for x in fila if x is not None]
        libro.close()
        assert not [c for c in celdas if c in ("nan", "inf", "-inf")
                    or c.startswith(("nan", "inf", "-inf"))], celdas


class WorkerFalso(QObject):
    """Un Worker que no lee nada ni arranca hilo ninguno."""
    progreso = Signal(int, int)
    terminado = Signal(object, str, str, object)
    gasto = Signal(str, float)
    fallo = Signal(str)
    creados = []

    def __init__(self, rutas, api_key):
        super().__init__()
        self.rutas = rutas
        self.fallos = []
        self._corriendo = False
        WorkerFalso.creados.append(self)

    def start(self):
        self._corriendo = True

    def isRunning(self):
        return self._corriendo


@pytest.fixture
def ventana(monkeypatch, tmp_path):
    from PySide6.QtWidgets import QMessageBox

    from facturas_excel.app import VentanaPrincipal
    # El aviso de las muestras (n.º 12) tiene sus propias pruebas más abajo.
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))
    WorkerFalso.creados = []
    monkeypatch.setattr(ventana_lectura, "Worker", WorkerFalso)
    monkeypatch.setattr(ventana_lectura, "leer_api_key", lambda: "clave-de-prueba")
    monkeypatch.setattr(sesion, "_ruta", lambda: str(tmp_path / "sesion.pkl.gz"))
    v = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=False)
    v._cola = []
    return v


def _dos_tacos_en_cola(v, tmp_path):
    """Dos documentos: el primero «leyéndose» y el segundo esperando turno."""
    rutas = []
    for i in range(2):
        ruta = tmp_path / f"taco {i + 1}.jpg"
        ruta.write_bytes(b"no se llega a leer")
        rutas.append(str(ruta))
    v.procesar_rutas(rutas)
    WorkerFalso.creados[0]._corriendo = False
    assert len(WorkerFalso.creados) == 1 and len(v._cola) == 1
    assert not v.btn_gastos.isEnabled()      # mientras se lee, Exportar apagado
    return rutas


def _recuperar(tmp_path):
    from facturas_excel.app import VentanaPrincipal
    return VentanaPrincipal(comprobar_updates=False, restaurar_sesion=True)


def _entregar(lectura_falsa, crudos):
    """El bloque leído llega como de verdad: por la señal de su lectura, y
    de ahí a la cola (ventana_lectura._incorporar)."""
    lectura_falsa._corriendo = False
    lectura_falsa.terminado.emit(procesar.preparar_lote(crudos, *CLIENTE),
                                 *CLIENTE, crudos)


def test_un_bloque_con_suplidos_nan_no_para_la_cola(ventana, tmp_path):
    _dos_tacos_en_cola(ventana, tmp_path)
    crudos = crudos_de(lectura(0, suplidos=NAN), lectura(1))
    _entregar(WorkerFalso.creados[0], crudos)

    assert len(WorkerFalso.creados) == 2          # la cola sigue con el 2.º
    assert ventana._cola_completados == 1
    # Exportar sigue apagado mientras se lee el 2.º (saldría un Excel a
    # medias) y vuelve en cuanto acaba la cola.
    assert not ventana.btn_gastos.isEnabled()
    _entregar(WorkerFalso.creados[1], crudos_de(lectura(2), origen="taco 2.jpg"))
    assert ventana.btn_gastos.isEnabled()
    assert ventana._timer_sesion.isActive()        # el guardado automático, armado
    assert not any(math.isnan(x) for fila in ventana.filas
                   for x in (fila.factura.base_iva, fila.factura.suplidos)
                   if isinstance(x, float))
    ventana._guardar_sesion()
    recuperada = _recuperar(tmp_path)
    assert recuperada.tabla.rowCount() == ventana.tabla.rowCount() > 0
    assert not sesion.apartada()


def test_una_sesion_vieja_con_nan_se_recupera_y_la_fila_sale_en_rojo(
        ventana, tmp_path):
    procesadas = lote(lectura(0), lectura(1))
    procesadas[0][1].facturas[0].base_iva = NAN     # como la dejaba la 1.25
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(procesadas, *CLIENTE)
    assert ventana.filas[0].estado == ERROR
    ventana._guardar_sesion()

    recuperada = _recuperar(tmp_path)
    assert recuperada.tabla.rowCount() == 2
    assert recuperada.filas[0].estado == ERROR
    assert not sesion.apartada()


def test_si_poner_un_bloque_falla_la_cola_sigue_y_queda_apuntado(
        ventana, tmp_path, monkeypatch):
    from facturas_excel import errores
    from facturas_excel.rutas import dir_datos

    _dos_tacos_en_cola(ventana, tmp_path)
    original = type(ventana)._revalidar_todo
    veces = []

    def revalidar_que_falla(self):
        veces.append(1)
        if len(veces) == 1:
            raise RuntimeError("fallo inventado al pintar el resumen")
        return original(self)
    monkeypatch.setattr(type(ventana), "_revalidar_todo", revalidar_que_falla)
    crudos = crudos_de(lectura(0), lectura(1))

    _entregar(WorkerFalso.creados[0], crudos)

    assert len(WorkerFalso.creados) == 2           # pasa al bloque siguiente
    assert ventana._cola_completados == 1
    # Con el 2.º leyéndose, Exportar no vuelve todavía (antes la red lo
    # encendía con bloques por leer).
    assert not ventana.btn_gastos.isEnabled()
    assert ventana._timer_sesion.isActive() and ventana._timer_muestras.isActive()
    with open(os.path.join(dir_datos(), errores.FICHERO), encoding="utf-8") as fh:
        assert "fallo inventado al pintar el resumen" in fh.read()
    _entregar(WorkerFalso.creados[1], crudos_de(lectura(2), origen="taco 2.jpg"))
    assert ventana._cola_completados == 2 and ventana.btn_gastos.isEnabled()


def test_un_fallo_despues_de_pasar_al_siguiente_no_arranca_otro(
        ventana, tmp_path, monkeypatch):
    _dos_tacos_en_cola(ventana, tmp_path)
    original = type(ventana)._iniciar_siguiente_cola

    veces = []

    def iniciar_y_fallar(self, *a, **k):
        original(self, *a, **k)
        veces.append(1)
        if len(veces) == 1:
            raise RuntimeError("fallo inventado al final")
    monkeypatch.setattr(type(ventana), "_iniciar_siguiente_cola", iniciar_y_fallar)
    crudos = crudos_de(lectura(0))

    _entregar(WorkerFalso.creados[0], crudos)

    assert len(WorkerFalso.creados) == 2           # uno solo más, no dos
    assert ventana._cola_completados == 1


# =====================================================================
# n.º 4 — Una hoja rara no tira el bloque; una lectura desbocada no tira
#          la otra, que era buena
# =====================================================================
def anidado(niveles: int):
    valor = []
    for _ in range(niveles):
        valor = [valor]
    return valor


@pytest.mark.parametrize("texto,esperado", [
    ('{"total": NaN, "suplidos": Infinity, "base_irpf": -Infinity}',
     {"total": None, "suplidos": None, "base_irpf": None}),
    ('{"total": 1e999, "pct_irpf": 15}', {"total": None, "pct_irpf": 15}),
    # Un entero de más de 15 cifras llega como texto, nunca como número (y
    # sin int(): con más de 4300 cifras lanzaba). Ver el del nº de factura.
    ('{"total": ' + "9" * 5000 + '}', {"total": "9" * 5000}),
    ('{"total": ' + "9" * 400 + '.5}', {"total": None}),
    ('{"total": 9999999999999999}', {"total": "9999999999999999"}),
], ids=["literales", "1e999", "entero_5000_cifras", "decimal_400_cifras",
        "entero_16_cifras"])
def test_el_json_de_gemini_sin_numeros_imposibles(texto, esperado):
    from facturas_excel.extraccion import _parse_json_tolerante
    assert _parse_json_tolerante(texto) == esperado


@pytest.mark.parametrize("texto", [
    "[" * 100_000 + "]" * 100_000,
    '{"concepto_texto": ' + "[" * 3000 + "]" * 3000 + "}",
    '{"total": 1' + "9" * 5000,                     # cortada en mitad del número
], ids=["100000_niveles", "3000_niveles_dentro", "cortada_en_un_numero"])
def test_el_json_desbocado_no_lanza(texto):
    from facturas_excel.extraccion import _parse_json_tolerante
    assert _parse_json_tolerante(texto) is None      # se vuelve a pedir


class ClienteGemini:
    """Contesta a cada modelo con un TEXTO fijo y apunta las llamadas."""

    def __init__(self, textos):
        from types import SimpleNamespace
        self.textos = textos
        self.llamadas = []
        self.models = SimpleNamespace(generate_content=self._generar)

    def _generar(self, model, contents, config):
        from types import SimpleNamespace
        self.llamadas.append(model)
        return SimpleNamespace(
            text=self.textos[model], model_version=model,
            usage_metadata=SimpleNamespace(prompt_token_count=1000,
                                           candidates_token_count=100,
                                           thoughts_token_count=20))


def extractor(monkeypatch, textos, modo="siempre"):
    from facturas_excel import extraccion
    monkeypatch.setattr(extraccion.genai, "Client", lambda **kw: None)
    ex = extraccion.Extractor("clave", modelos=["modelo-a", "modelo-b"],
                              modo_doble=modo)
    ex.client = ClienteGemini(textos)
    return ex


@pytest.mark.parametrize("desbocada", [
    "[" * 100_000 + "]" * 100_000,
    '{"total": ' + "9" * 5000 + ', "lineas_iva": [',
    '{"total": 1' + "9" * 5000,
], ids=["100000_niveles", "entero_5000_cifras_cortada", "cortada_en_un_numero"])
def test_una_lectura_desbocada_no_tira_la_otra(monkeypatch, desbocada):
    buena = json.dumps(lectura(0))
    ex = extractor(monkeypatch, {"modelo-a": desbocada, "modelo-b": buena})

    leido = ex.extraer(b"img", "taco.pdf", 1)

    assert leido.crudo["num_factura"] == "F26/00100"   # la del otro modelo
    assert leido.crudo.get("_error_2")                 # y se dice que la 1.ª falló
    assert ex.client.llamadas.count("modelo-a") == 3   # se reintentó, se pagó
    assert len(leido.consumos) == 4


class RespuestaRota:
    """Una respuesta de la que ni se puede sacar el texto (un fallo que no se
    ha previsto en el SDK o en la respuesta)."""
    model_version = "modelo-roto"
    usage_metadata = None

    @property
    def text(self):
        raise RuntimeError("respuesta sin texto que leer")


@pytest.mark.parametrize("modo,rota,buena", [
    ("siempre", "modelo-a", "modelo-b"),
    ("siempre", "modelo-b", "modelo-a"),
    ("dudosas", "modelo-b", "modelo-a"),
], ids=["siempre_la_primera", "siempre_la_segunda", "dudosas_la_segunda"])
def test_un_fallo_raro_de_una_lectura_no_tira_la_otra(monkeypatch, modo, rota, buena):
    # Confianza media: en «dudosas» también se pide la segunda.
    ex = extractor(monkeypatch, {buena: json.dumps(lectura(0, confianza="media"))},
                   modo=modo)
    generar = ex.client._generar

    def generar_o_romper(model, contents, config):
        if model == rota:
            ex.client.llamadas.append(model)
            return RespuestaRota()
        return generar(model, contents, config)
    ex.client.models.generate_content = generar_o_romper

    leido = ex.extraer(b"img", "taco.pdf", 1)

    assert leido.crudo["num_factura"] == "F26/00100"       # la buena se queda
    assert "respuesta sin texto" in leido.crudo["_error_2"]  # y se dice qué pasó
    assert rota in ex.client.llamadas


def test_una_lectura_con_un_entero_enorme_se_queda(monkeypatch):
    rara = json.dumps(lectura(0))[:-1] + ', "suplidos": ' + "9" * 5000 + "}"
    ex = extractor(monkeypatch, {"modelo-a": rara, "modelo-b": rara}, modo="no")
    leido = ex.extraer(b"img", "taco.pdf", 1)
    assert leido.crudo["suplidos"] is None and leido.crudo["total"] == 60.5
    assert ex.client.llamadas == ["modelo-a"]


def test_un_numero_de_factura_largo_escrito_sin_comillas_se_queda(monkeypatch):
    # Sin esquema, el nº de factura puede llegar como número, y los de las
    # eléctricas o las telefónicas tienen 16-20 cifras: pasaba a null sin
    # avisar y la factura se quedaba sin número.
    texto = json.dumps(lectura(0, num_factura=0, suplidos=0)).replace(
        '"num_factura": 0', '"num_factura": 21240000012345678').replace(
        '"suplidos": 0', '"suplidos": 9999999999999999')
    ex = extractor(monkeypatch, {"modelo-a": texto}, modo="no")
    leido = ex.extraer(b"img", "taco.pdf", 1)
    assert leido.crudo["num_factura"] == "21240000012345678"
    assert lote(leido.crudo)[0][1].facturas[0].num_factura == "21240000012345678"
    # Donde va un importe sigue sin valer, pero ahora se dice (ámbar).
    assert leido.crudo["suplidos"] is None and leido.crudo.get("_saneado")


RARAS = {
    "lineas_texto": dict(lineas_iva="21%"),
    "lineas_con_none": dict(lineas_iva=[None]),
    "lineas_con_numero": dict(lineas_iva=[5]),
    "lineas_numero": dict(lineas_iva=5),
    "lineas_dict": dict(lineas_iva={"base": 10}),
    "ultima_pagina_texto": dict(_ultima_pagina_consolidada="abc"),
    "union_manual_rara": dict(_paginas_union_manual=[[1, 2, 3]]),
    "anidado_900": dict(concepto_texto=anidado(900)),
    "nombre_lista": dict(emisor_nombre=["PROVEEDOR", "UNO"]),
    "nombre_dict": dict(emisor_nombre={"nombre": "X"}),
}


@pytest.mark.parametrize("cambios", list(RARAS.values()), ids=list(RARAS))
def test_una_hoja_rara_no_tira_el_bloque_de_25(cambios):
    lecturas = [lectura(i) for i in range(25)]
    lecturas[6] = lectura(6, **cambios)

    procesadas = procesar.preparar_lote(crudos_de(*lecturas), *CLIENTE)

    assert len(procesadas) == 25
    assert sorted(pr.pagina for _, pr in procesadas) == list(range(1, 26))
    rara = next(pr for _, pr in procesadas if pr.pagina == 7)
    otras = [pr for _, pr in procesadas if pr.pagina != 7]
    assert not any("HOJA NO LEÍDA" in pr.aviso for pr in otras)
    assert all(pr.facturas[0].num_factura for pr in otras)
    if "HOJA NO LEÍDA" in rara.aviso:            # la rara, en rojo con su motivo
        assert validar(rara.facturas[0]).estado == ERROR


@pytest.mark.parametrize("nombre", [["PROVEEDOR", "UNO"], {"nombre": "X"}, 12345],
                         ids=["lista", "dict", "numero"])
def test_un_nombre_raro_no_impide_saber_el_cliente(nombre):
    lecturas = [lectura(i) for i in range(25)]
    lecturas[6] = lectura(6, emisor_nombre=nombre)
    lecturas[7] = lectura(7, receptor_nombre=nombre, receptor_nif=["12345678Z"])
    assert procesar.detectar_cliente(lecturas) == CLIENTE
    # También la ventana, que analiza el lote entero tras cada bloque.
    assert procesar.analizar_cliente(lecturas).mejor.nif == CLIENTE[1]


def test_si_una_hoja_hace_fallar_la_busqueda_del_cliente_se_sigue_sin_ella(
        monkeypatch):
    # Un fallo que no se ha previsto (aquí, inventado) en una sola hoja.
    limpiar = procesar.limpiar

    def limpiar_que_falla(texto):
        if texto == "NOMBRE QUE FALLA":
            raise RuntimeError("fallo inventado")
        return limpiar(texto)
    monkeypatch.setattr(procesar, "limpiar", limpiar_que_falla)
    lecturas = [lectura(i) for i in range(25)]
    lecturas[6] = lectura(6, emisor_nombre="NOMBRE QUE FALLA")
    assert procesar.detectar_cliente(lecturas) == CLIENTE


@pytest.mark.parametrize("lineas", ["21%", 5, {"base": 10}, None],
                         ids=["texto", "numero", "dict", "nada"])
def test_comparar_dos_lecturas_con_lineas_raras_no_lanza(lineas):
    from facturas_excel.doble_lectura import comparar, es_dudosa
    rara = lectura(0, lineas_iva=lineas)
    diferencias = comparar(lectura(0), rara)
    assert [d["campo"] for d in diferencias] == ["lineas_iva"]
    assert es_dudosa(rara)                  # sin desglose, que la mire otro


# =====================================================================
# n.º 14 — La lectura se sanea a la entrada de la doble lectura: solo las
#           claves del esquema y las internas conocidas, con su tipo y tamaño
# =====================================================================
def combinada(principal, segunda=None):
    from facturas_excel.doble_lectura import combinar
    if segunda is None:
        return combinar(principal, None, "modelo-a", "")
    return combinar(principal, segunda, "modelo-a", "modelo-b")


def test_una_lectura_normal_pasa_igual():
    d = lectura(0)
    c = combinada(d, copy.deepcopy(d))
    assert {k: c[k] for k in d} == d
    assert "_saneado" not in c and c["_discrepancias"] == []
    assert "Gemini traía datos" not in lote(d)[0][1].aviso


@pytest.mark.parametrize("campo,valor,esperado", [
    ("emisor_nombre", ["PROVEEDOR", "UNO"], None),
    ("emisor_nombre", {"nombre": "X"}, None),
    ("emisor_nombre", 12345, "12345"),
    ("emisor_nif", True, None),
    ("num_factura", 10 ** 400, None),
    ("num_factura", NAN, None),
    ("fecha", ["01/01/2026"], None),
    ("total", "abc", None),
    ("total", [1, 2], None),
    ("total", "9" * 5000, None),
    ("suplidos", NAN, None),
    ("es_bien_inversion", "false", False),
    ("manuscrito_en_importes", [], False),
    ("tipo_documento", "abono_raro", None),
    ("tipo_documento", ["factura"], None),
    ("confianza", 5, None),
    ("estado_pagina_factura", 5, None),
    ("lineas_iva", "21%", []),
    ("lineas_iva", 5, []),
    ("lineas_iva", [None, {"base": "x", "tipo_iva": 21, "cuota_iva": NAN}],
     [{"base": None, "tipo_iva": 21.0, "cuota_iva": None, "pct_requiv": None,
       "cuota_requiv": None}]),
], ids=lambda v: corto(v) if not isinstance(v, str) else v[:20])
def test_cada_dato_queda_con_su_tipo(campo, valor, esperado):
    c = combinada(lectura(0, **{campo: valor}))
    assert c[campo] == esperado
    # Un número por nombre o «false» en texto se entienden sin más; lo
    # demás se apunta para avisar en ámbar.
    if (campo, valor) not in (("emisor_nombre", 12345), ("es_bien_inversion", "false")):
        assert c.get("_saneado")


def test_los_textos_desbocados_se_recortan():
    c = combinada(lectura(0, emisor_nombre="PROVEEDOR LARGO " * 625,
                          num_factura="9" * 5000, moneda="x" * 1000))
    assert len(c["emisor_nombre"]) <= 300 and len(c["num_factura"]) <= 200
    assert len(c["moneda"]) <= 20
    pr = lote(c)[0][1]
    assert len(pr.facturas[0].num_factura) <= 60    # y la fila, como siempre


def test_lo_saneado_sale_en_ambar_diciendo_que():
    c = combinada(lectura(0, tipo_documento="abono_raro", confianza=5))
    pr = lote(c)[0][1]
    assert "tipo de documento «abono_raro»" in pr.aviso
    assert "confianza" in pr.aviso


@pytest.mark.parametrize("clave,valor", [
    ("_ultima_pagina_consolidada", 10 ** 9), ("_ultima_pagina_consolidada", "abc"),
    ("_ultima_pagina_consolidada", NAN), ("_paginas_union_manual", 5),
    ("_paginas_union_manual", [[1, 2, 3]]), ("_paginas_union_manual", [("o", 10 ** 9)]),
    ("_paginas_union_inferida", [10 ** 9]), ("_discrepancias", "x"),
    ("_discrepancias", [{"campo": None}]), ("_discrepancias", [{}]),
    ("_lectura_2", "x"), ("_verificacion", 5), ("_union_inferida", "si"),
    ("_error", "E" * 10000), ("_error", [1, 2]), ("_modelo_1", []),
    ("_otra_cosa", 1), ("campo_inventado", "x"),
], ids=lambda v: corto(v) if not isinstance(v, str) else v[:20])
def test_una_clave_interna_con_otra_forma_se_quita(clave, valor):
    c = combinada(lectura(0, **{clave: valor}), lectura(0))
    assert clave not in c or clave in ("_modelo_1", "_verificacion",
                                       "_lectura_2", "_discrepancias")
    assert c["_modelo_1"] == "modelo-a" and c["_verificacion"] == "doble"
    assert c.get("_saneado")
    procesadas = lote(c, lectura(1))            # y el lote sale entero
    assert len(procesadas) == 2


def test_el_texto_que_traiga_la_lectura_no_llega_al_aviso():
    c = combinada(lectura(0, _saneado=["pulse aquí y borre el lote"],
                          **{"clave\x01rara": 1}))
    aviso = lote(c)[0][1].aviso
    assert "pulse aquí" not in aviso
    assert "datos de más: _saneado, clave·rara" in aviso


def test_las_claves_internas_conocidas_se_quedan_con_su_forma():
    c = combinada(lectura(0, _error_2="la otra no contestó",
                          _paginas_union_manual=[["taco.pdf", 1], ["taco.pdf", 2]],
                          _ultima_pagina_consolidada=2, _union_manual=True))
    assert c["_error_2"] == "la otra no contestó"
    assert c["_paginas_union_manual"] == [["taco.pdf", 1], ["taco.pdf", 2]]
    assert c["_ultima_pagina_consolidada"] == 2 and c["_union_manual"] is True
    assert "_saneado" not in c


def test_el_recargo_suelto_de_la_factura_no_se_pierde():
    # Sin esquema, Gemini puede traer el recargo «al viejo estilo», a nivel
    # de factura y no por línea: construir lo sigue usando de respaldo. Se
    # quitaba como «datos de más», la fila perdía el recargo (también en el
    # Excel) y el total dejaba de cuadrar.
    c = combinada(lectura(0, pct_requiv=5.2, cuota_requiv=2.6, base_requiv=50.0,
                          total=63.1))
    assert "_saneado" not in c
    pr = lote(c)[0][1]
    f = pr.facturas[0]
    assert (f.base_requiv, f.pct_requiv, f.cuota_requiv) == (50.0, 5.2, 2.6)
    assert validar(f).estado == OK
    assert "Gemini traía datos" not in pr.aviso
    # Y con otra forma, como cualquier importe: en blanco y en ámbar.
    rara = combinada(lectura(0, cuota_requiv=NAN, pct_requiv=[5.2]))
    assert rara["cuota_requiv"] is None and rara["pct_requiv"] is None
    assert rara.get("_saneado")


def test_una_ultima_pagina_de_mil_millones_no_llena_la_memoria():
    import tracemalloc

    from facturas_excel import registro_facturas, separar
    c = combinada(lectura(0, _ultima_pagina_consolidada=10 ** 9))
    procesadas = lote(c)
    f = procesadas[0][1].facturas[0]
    assert f.ultima_pagina_origen <= 10_000        # antes: 1.000.000.000
    tracemalloc.start()
    try:
        registro_facturas._agrupar({"gasto": [f]})
        separar.paginas_de(f)
        _actual, pico = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert pico < 20 * 2 ** 20


def test_modo_dudosas_no_paga_otra_lectura_por_un_false_en_texto(monkeypatch):
    buena = json.dumps(lectura(0, manuscrito_en_importes="false"))
    ex = extractor(monkeypatch, {"modelo-a": buena, "modelo-b": buena},
                   modo="dudosas")
    leido = ex.extraer(b"img", "taco.pdf", 1)
    assert leido.crudo["manuscrito_en_importes"] is False
    assert ex.client.llamadas == ["modelo-a"]          # sin segunda lectura


# =====================================================================
# n.º 12 — Muestras con \ud800 o NaN; n.º 13 — localizar con enteros enormes
# =====================================================================
@pytest.mark.parametrize("valor", ["PROVEEDOR \ud800 SL", "\udfff" * 3, NAN, INF,
                                   10 ** 5000, {"a": [NAN, "\ud800"]}],
                         ids=["surrogate", "surrogates", "nan", "inf",
                              "entero_5000_cifras", "anidado"])
def test_las_muestras_se_guardan_con_cualquier_lectura(valor):
    from facturas_excel import muestras_revision
    contenido = muestras_revision._json({"datos": {"emisor_nombre": valor}})
    assert isinstance(json.loads(contenido.decode("utf-8")), dict)
    muestras_revision.guardar_lecturas(
        [(b"", "taco.pdf", 1, {"emisor_nombre": valor, "total": NAN})])


def test_un_bloque_con_letras_sueltas_no_avisa_de_las_muestras(
        ventana, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    avisos = []
    monkeypatch.setattr(QMessageBox, "warning",
                        staticmethod(lambda *a, **k: avisos.append(a[1:3])))
    crudos = crudos_de(combinada(lectura(0, emisor_nombre="PROVEEDOR \ud800 SL")),
                       combinada(lectura(1, num_factura="\ud800123")))
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(procesar.preparar_lote(crudos, *CLIENTE), *CLIENTE, crudos)
    ventana._guardar_muestra_revision()
    assert ventana.tabla.rowCount() == 2
    assert avisos == []


@pytest.mark.parametrize("valor", [10 ** 400, -10 ** 400, 10 ** 5000, NAN, INF,
                                   "9" * 5000],
                         ids=["10e400", "-10e400", "10e5000", "nan", "inf",
                              "texto_5000_cifras"])
def test_localizar_compara_cualquier_valor(valor):
    from facturas_excel import localizar
    assert isinstance(localizar._comparable(valor), str)
    caja = localizar.Caja("total", "121,00", 0.1, 0.1, 0.2, 0.2)
    assert localizar.cajas_de([caja], "total", valor) == []


# =====================================================================
# n.º 25 — Muchísimas líneas de IVA: se juntan por tipo y la ficha no se
#           congela
# =====================================================================
def lineas(n, base=1.0, tipo=21.0, recargo=None):
    return [{"base": base, "tipo_iva": tipo, "cuota_iva": round(base * tipo / 100, 2),
             "pct_requiv": recargo,
             "cuota_requiv": None if recargo is None else round(base * recargo / 100, 2)}
            for _ in range(n)]


@pytest.mark.parametrize("n", [13, 60, 300, 3000])
def test_muchas_lineas_del_mismo_tipo_quedan_en_una(n):
    c = combinada(lectura(0, lineas_iva=lineas(n), total=round(n * 1.21, 2)))
    assert c["lineas_iva"] == [{"base": float(n), "tipo_iva": 21.0,
                                "cuota_iva": round(n * 0.21, 2), "pct_requiv": None,
                                "cuota_requiv": None, "_juntadas": n,
                                "_cuota_lineas": round(n * 0.21, 2),
                                "_requiv_lineas": None, "_lineas_ambar": 0}]
    pr = lote(c)[0][1]
    assert len(pr.facturas) == 1
    assert f"{n} líneas de IVA" in pr.aviso        # en ámbar, diciendo qué
    assert validar(pr.facturas[0]).estado == OK     # y cuadra


@pytest.mark.parametrize("n, base, tipo, recargo", [
    (30, 15.40, 4.0, 0.5),     # resumen de 30 albaranes de pan a un bar
    (13, 10.10, 21.0, 5.2),
], ids=["30_albaranes_al_4", "13_lineas_al_21"])
def test_lineas_juntadas_que_redondean_cuadran(n, base, tipo, recargo):
    # Cada línea, con su IVA y su recargo redondeados a céntimos, cuadra sola
    # (en la 1.25 salían n filas en verde). Juntadas, la suma de las cuotas
    # se aparta de base_total×% hasta medio céntimo por línea: en rojo no se
    # podía «Marcar revisada» y había que teclear una cuota que no es la
    # cobrada (30 × 15,40: IVA 18,60 frente a 18,48 y recargo 2,40 frente a
    # 2,31). El redondeo de cada línea se apunta al leer, y con él cuadra.
    juntas = lineas(n, base, tipo, recargo)
    total = round(sum(x["base"] + x["cuota_iva"] + x["cuota_requiv"]
                      for x in juntas), 2)
    for segunda in (None, lectura(0, lineas_iva=lineas(n, base, tipo, recargo),
                                  total=total)):
        c = combinada(lectura(0, lineas_iva=juntas, total=total), segunda)
        f = lote(c)[0][1].facturas[0]
        assert (f.cuota_iva, f.cuota_requiv) == (
            round(n * round(base * tipo / 100, 2), 2),
            round(n * round(base * recargo / 100, 2), 2))   # las cobradas
        assert f.lineas_juntadas == n
        resultado = validar(f)
        assert resultado.estado == OK, resultado.mensajes


def test_la_fila_juntada_se_puede_marcar_revisada_y_exportar(ventana):
    # 30 albaranes de 15,40 € al 4 %: IVA cobrado 18,60, base×% = 18,48.
    juntas = lineas(30, 15.40, 4.0)
    procesadas = lote(combinada(lectura(0, lineas_iva=juntas, total=480.6)))
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(procesadas, *CLIENTE)
    assert len(ventana.filas) == 1 and ventana.filas[0]["estado"] == REVISAR
    assert ventana._marcar_revisada([0]) is not None
    assert not ventana.filas[0].pendiente
    assert ventana.filas[0].factura.cuota_iva == 18.6      # la cobrada


def _elegir_lectura_2_del_desglose(ventana, principal, segunda, total):
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(lote(combinada(lectura(0, lineas_iva=principal, total=total),
                                         lectura(0, lineas_iva=segunda, total=total))),
                          *CLIENTE)
    f = ventana.filas[0].factura
    indice = next(i for i, d in enumerate(f.discrepancias)
                  if d.get("campo") == "lineas_iva")
    ventana._resolver_discrepancia(0, indice, 2)
    ventana._revalidar_todo()
    return ventana.filas[0].factura, [str(m) for m in ventana.filas[0].get("mensajes") or []]


UNA_LINEA_DE_PAN = [{"base": 462.0, "tipo_iva": 4.0, "cuota_iva": 18.48,
                     "pct_requiv": None, "cuota_requiv": None}]


def test_usar_la_lectura_2_juntada_se_lleva_su_redondeo(ventana):
    # La principal leyó una línea con la cuota calculada; la segunda, los 30
    # albaranes (18,60, la cobrada). Al elegir la segunda, la fila es su
    # suma: cuadra con el redondeo de cada albarán, no en rojo sin salida.
    f, mensajes = _elegir_lectura_2_del_desglose(
        ventana, UNA_LINEA_DE_PAN, lineas(30, 15.40, 4.0), 480.6)
    assert (f.base_iva, f.cuota_iva) == (462.0, 18.6)
    assert not any("Cuota IVA descuadra" in m for m in mensajes), mensajes
    assert ventana.filas[0]["estado"] != ERROR
    assert f.lineas_juntadas == 30


def test_usar_la_lectura_2_de_una_linea_olvida_las_juntadas(ventana):
    # Al revés: la principal juntó 30 y la segunda leyó una. Tras elegirla,
    # la fila ya no es una suma: una cuota mal tecleada después no se
    # disculpa con el redondeo de los 30 albaranes.
    f, _ = _elegir_lectura_2_del_desglose(
        ventana, lineas(30, 15.40, 4.0), UNA_LINEA_DE_PAN, 480.6)
    assert f.cuota_iva == 18.48
    f.cuota_iva = 18.63
    ventana._revalidar_todo()
    assert ventana.filas[0]["estado"] == ERROR, ventana.filas[0].get("mensajes")
    assert (f.lineas_juntadas, f.redondeo_lineas_iva) == (1, 0.0)


def test_una_cuota_mal_leida_entre_las_juntadas_sigue_en_rojo():
    # El margen es el del redondeo, no más: un euro de más en una línea no
    # se disculpa por estar juntada.
    juntas = lineas(13, 10.10, 21.0, 5.2)
    juntas[5] = dict(juntas[5], cuota_iva=3.12, cuota_requiv=1.53)
    f = lote(combinada(lectura(0, lineas_iva=juntas)))[0][1].facturas[0]
    mensajes = validar(f).mensajes
    assert validar(f).estado == ERROR
    assert any("Cuota IVA descuadra" in m for m in mensajes)
    assert any("Cuota del recargo descuadra" in m for m in mensajes)


@pytest.mark.parametrize("n, error", [(60, 0.30), (200, 1.00)])
def test_una_cuota_mal_leida_entre_muchas_juntadas_sigue_en_rojo(n, error):
    # 60 líneas de 10,00 € al 21 % + 5,2 %: cada cuota es justa, el redondeo
    # no aparta nada. Una mal leída (2,40 por 2,10) con el total impreso
    # bueno: en la 1.25 esa línea iba en rojo. Con un margen de medio
    # céntimo por línea (0,31 € con 60, 1,01 € con 200) pasaba en ámbar
    # como «redondeo» y se exportaba con «Marcar revisada».
    juntas = lineas(n, 10.0, 21.0, 5.2)
    total = round(sum(x["base"] + x["cuota_iva"] + x["cuota_requiv"]
                      for x in juntas), 2)
    juntas[7] = dict(juntas[7], cuota_iva=round(2.10 + error, 2),
                     cuota_requiv=round(0.52 + error, 2))
    f = lote(combinada(lectura(0, lineas_iva=juntas, total=total)))[0][1].facturas[0]
    assert f.lineas_juntadas == n
    resultado = validar(f)
    mensajes = [str(m) for m in resultado.mensajes]
    assert resultado.estado == ERROR, mensajes
    assert any("Cuota IVA descuadra" in m for m in mensajes)
    assert any("Cuota del recargo descuadra" in m for m in mensajes)
    assert not any("redondeo" in m for m in mensajes)


def test_el_redondeo_de_las_juntadas_cuenta_solo_lo_que_aparta_cada_linea():
    # 30 albaranes de 15,40 € al 4 %: cobrado 18,60, base×% 18,48. Con una
    # línea leída alta, la fila queda como habría quedado esa línea sola en
    # la 1.25: 0,01 cuadra, 0,04 en ámbar (redondeo) y 0,10 en rojo.
    for de_mas, estado in ((0.01, OK), (0.04, REVISAR), (0.10, ERROR)):
        juntas = lineas(30, 15.40, 4.0)
        juntas[3] = dict(juntas[3], cuota_iva=round(0.62 + de_mas, 2))
        total = round(480.6 + de_mas, 2)
        f = lote(combinada(lectura(0, lineas_iva=juntas, total=total)))[0][1].facturas[0]
        assert validar(f).estado == estado, (de_mas, validar(f).mensajes)


@pytest.mark.parametrize("signo_lineas", [1, -1], ids=["lineas_en_positivo",
                                                       "lineas_en_negativo"])
def test_un_abono_de_lineas_juntadas_tambien_cuadra(signo_lineas):
    # El abono de los 30 albaranes: el total en negativo. Si Gemini dejó las
    # líneas en positivo, procesar las pasa a negativo; el redondeo de cada
    # una va en el sentido de la base y sigue valiendo.
    juntas = [{k: (signo_lineas * v if k in ("base", "cuota_iva", "cuota_requiv")
                   and v is not None else v) for k, v in x.items()}
              for x in lineas(30, 15.40, 4.0, 0.5)]
    f = lote(combinada(lectura(0, lineas_iva=juntas, total=-483.0)))[0][1].facturas[0]
    assert (f.base_iva, f.cuota_iva, f.cuota_requiv) == (-462.0, -18.6, -2.4)
    assert validar(f).estado == OK, validar(f).mensajes


def _pan_en_la_tabla(ventana, juntas=None, total=480.6):
    """Los 30 albaranes de pan, juntados en una fila de la tabla."""
    juntas = lineas(30, 15.40, 4.0) if juntas is None else juntas
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(lote(combinada(lectura(0, lineas_iva=juntas, total=total))),
                          *CLIENTE)
    return ventana.filas[0].factura


def _teclear(ventana, columna, texto):
    ventana.tabla.item(0, columna).setText(texto)
    return [str(m) for m in ventana.filas[0].get("mensajes") or []]


def test_cambiar_el_tipo_de_la_fila_juntada_no_deja_buena_la_cuota_vieja(ventana):
    # El redondeo de las líneas es para su base y su tipo: si a mano se pone
    # otro tipo y no se cambia la cuota, no cuadra. Pasa por la tabla: lo que
    # teclea una persona ya no es la suma leída (ver la de abajo).
    from facturas_excel.tabla_facturas import C_CUOTA, C_PCT, C_TOTAL
    f = _pan_en_la_tabla(ventana)
    assert validar(f).estado == OK
    _teclear(ventana, C_TOTAL, "508,20")
    mensajes = _teclear(ventana, C_PCT, "10")
    assert any("Cuota IVA descuadra" in m for m in mensajes), mensajes
    assert ventana.filas[0]["estado"] == ERROR
    mensajes = _teclear(ventana, C_CUOTA, "46,20")   # 462 × 10 %: base×%, buena
    assert validar(f).estado == OK, mensajes


def test_el_tipo_corregido_a_mano_da_por_buenas_las_cuotas_leidas(ventana):
    # Los 30 albaranes al 10 % (1,54 cada uno: 46,20) con el tipo leído 4 %:
    # en rojo. Al corregir el tipo, la cuota leída es la buena; el redondeo
    # apuntado era el del 4 % y no puede dejarla en rojo (46,20 frente a
    # 46,32), que obligaría a teclear la misma cuota que ya está.
    from facturas_excel.tabla_facturas import C_PCT
    juntas = [dict(x, cuota_iva=1.54) for x in lineas(30, 15.40, 4.0)]
    f = _pan_en_la_tabla(ventana, juntas, total=508.2)
    assert validar(f).estado == ERROR
    mensajes = _teclear(ventana, C_PCT, "10")
    assert validar(f).estado == OK, mensajes


@pytest.mark.parametrize("n, base, linea, leida", [
    (60, 100.0, 7, 21.30),      # una diezmilésima de 6.000 €: 0,60
    (40, 500.0, 3, 106.00),     # una diezmilésima de 20.000 €: 2,00
], ids=["60_de_100_con_0_30_de_mas", "40_de_500_con_1_euro_de_mas"])
def test_una_cuota_mal_leida_entre_lineas_grandes_juntadas_sigue_en_rojo(
        n, base, linea, leida):
    # Cada cuota es justa (sin redondeo) y una se lee de más, con el total
    # impreso bueno. En la 1.25 esa línea iba en rojo: su ámbar «redondeo
    # por líneas» es el de SU base (0,05). Juntadas, ese ámbar se medía con
    # la base de toda la fila y la cuota de más salía en ámbar y se exportaba
    # con «Marcar revisada», aunque el redondeo de cada línea se sabía: cero.
    juntas = lineas(n, base, 21.0)
    total = round(sum(x["base"] + x["cuota_iva"] for x in juntas), 2)
    juntas[linea] = dict(juntas[linea], cuota_iva=leida)
    f = lote(combinada(lectura(0, lineas_iva=juntas, total=total)))[0][1].facturas[0]
    assert f.lineas_juntadas == n
    resultado = validar(f)
    mensajes = [str(m) for m in resultado.mensajes]
    assert resultado.estado == ERROR, mensajes
    assert any("Cuota IVA descuadra" in m for m in mensajes)
    assert not any("redondeo" in m for m in mensajes)


def test_lineas_grandes_juntadas_con_su_redondeo_salen_en_ambar_como_en_la_1_25():
    # Lo que sí es redondeo: tres líneas de 2.000 € cuya cuota se aparta 0,10
    # de base×% (en la 1.25, cada una en ámbar: su margen es 0,20). Juntadas
    # se apartan 0,30 en total: ámbar, no rojo sin salida.
    juntas = lineas(13, 2000.0, 21.0)
    for i in (2, 4, 6):
        juntas[i] = dict(juntas[i], cuota_iva=420.10)
    total = round(sum(x["base"] + x["cuota_iva"] for x in juntas), 2)
    f = lote(combinada(lectura(0, lineas_iva=juntas, total=total)))[0][1].facturas[0]
    resultado = validar(f)
    mensajes = [str(m) for m in resultado.mensajes]
    assert resultado.estado == REVISAR, mensajes
    assert any("redondeo" in m for m in mensajes)


@pytest.mark.parametrize("n, base, recargo, linea, cambio", [
    (30, 15.40, 0.5, 3, {"cuota_iva": 0.51}),      # 0,62 leída 0,51
    (30, 15.40, 0.5, 3, {"cuota_iva": 0.52}),
    (30, 15.40, 0.5, 5, {"cuota_requiv": None}),   # el recargo, perdido
    (30, 15.40, 0.5, 5, {"cuota_requiv": 0.0}),
    (200, 1.10, None, 50, {"cuota_iva": 0.43}),    # 0,23 leída 0,43
], ids=["pan_iva_0_51", "pan_iva_0_52", "pan_sin_recargo", "pan_recargo_0",
        "200_lineas_iva_0_43"])
def test_una_cuota_mal_leida_que_compensa_el_redondeo_sigue_en_rojo(
        n, base, recargo, linea, cambio):
    # Los 30 albaranes de pan al 4 % + 0,5 %: IVA cobrado 18,60 (base×%
    # 18,48) y recargo 2,40 (2,31). Una cuota leída de menos acerca la suma
    # a base×% (18,49; 2,32) y cuadraba con ella: «la que más se acerque».
    # Con 200 líneas de 1,10 € (0,23 cada una: 46,00; base×% 46,20), una
    # leída de más. En la 1.25 esa línea iba en rojo; aquí solo quedaba el
    # ámbar del total y con «Marcar revisada» se exportaba la cuota mal
    # leída. La cuota leída se compara solo con la suma de sus líneas.
    juntas = lineas(n, base, 4.0 if recargo else 21.0, recargo)
    total = round(sum(x["base"] + x["cuota_iva"] + (x["cuota_requiv"] or 0)
                      for x in juntas), 2)
    juntas[linea] = dict(juntas[linea], **cambio)
    f = lote(combinada(lectura(0, lineas_iva=juntas, total=total)))[0][1].facturas[0]
    assert f.lineas_juntadas == n
    resultado = validar(f)
    assert resultado.estado == ERROR, [str(m) for m in resultado.mensajes]


def test_la_cuota_tecleada_en_la_fila_juntada_puede_ser_la_cobrada(ventana):
    # Lo que teclea una persona sí se compara con las dos: la suma de las
    # líneas (18,60, lo cobrado) o base×% (18,48). Así, la cuota mal leída
    # en rojo se corrige tecleando la cobrada…
    from facturas_excel.tabla_facturas import C_CUOTA
    juntas = lineas(30, 15.40, 4.0)
    juntas[3] = dict(juntas[3], cuota_iva=0.51)
    f = _pan_en_la_tabla(ventana, juntas)
    assert validar(f).estado == ERROR
    mensajes = _teclear(ventana, C_CUOTA, "18,60")
    assert validar(f).estado == OK, mensajes


def test_la_cuota_tecleada_en_la_fila_juntada_puede_ser_base_por_tipo(ventana):
    # … y la de una factura que, con los albaranes desglosados, cobra el IVA
    # sobre el total (480,48), tecleando la de base×%.
    from facturas_excel.tabla_facturas import C_CUOTA
    f = _pan_en_la_tabla(ventana, total=480.48)
    assert any("total no cuadra" in str(m) for m in validar(f).mensajes)
    mensajes = _teclear(ventana, C_CUOTA, "18,48")
    assert validar(f).estado == OK, mensajes


def test_se_juntan_por_tipo_de_iva_y_de_recargo():
    juntas = (lineas(10, 10.0, 21.0, 5.2) + lineas(10, 1.0, 21.0)
              + lineas(10, 2.0, 10.0, 1.4))
    c = combinada(lectura(0, lineas_iva=juntas))
    claves = [(x["tipo_iva"], x["pct_requiv"], x["base"]) for x in c["lineas_iva"]]
    assert claves == [(21.0, 5.2, 100.0), (21.0, None, 10.0), (10.0, 1.4, 20.0)]
    assert c["lineas_iva"][0]["cuota_requiv"] == 5.2


def test_juntar_lineas_con_tres_decimales_no_descuadra():
    # Importes de artículos con tres decimales (carburante, granel…): con
    # cada suma redondeada a céntimos el error se acumulaba línea a línea
    # (13 líneas de 1,005 € daban 13,02 en vez de 13,065).
    raras = [{"base": 1.005, "tipo_iva": 21.0, "cuota_iva": 0.211, "pct_requiv": 5.2,
              "cuota_requiv": 0.052}] * 13
    raras += [{"base": round(0.5 + i * 0.731, 3), "tipo_iva": 10.0,
               "cuota_iva": round((0.5 + i * 0.731) * 0.1, 3), "pct_requiv": None,
               "cuota_requiv": None} for i in range(60)]
    c = combinada(lectura(0, lineas_iva=raras))
    for tipo, juntas in ((21.0, c["lineas_iva"][0]), (10.0, c["lineas_iva"][1])):
        de_ese_tipo = [x for x in raras if x["tipo_iva"] == tipo]
        for campo in ("base", "cuota_iva", "cuota_requiv"):
            leido = math.fsum(x[campo] or 0 for x in de_ese_tipo)
            assert abs((juntas[campo] or 0) - leido) < 1e-6, (tipo, campo)


def test_con_tipos_inventados_se_quedan_doce_como_mucho():
    inventadas = [dict(linea, tipo_iva=float(i)) for i, linea in enumerate(lineas(300))]
    c = combinada(lectura(0, lineas_iva=inventadas))
    assert len(c["lineas_iva"]) == 12
    assert len(lote(c)[0][1].facturas) == 12
    assert any("quitadas" in nota for nota in c["_saneado"])


def test_hasta_doce_lineas_todo_sigue_como_antes():
    c = combinada(lectura(0, lineas_iva=lineas(12)))
    assert len(c["lineas_iva"]) == 12 and "_saneado" not in c


def test_la_segunda_lectura_con_300_lineas_tambien_se_junta():
    c = combinada(lectura(0, lineas_iva=lineas(300), total=363.0),
                  lectura(0, lineas_iva=lineas(300), total=363.0))
    assert c["_discrepancias"] == []


def test_la_ficha_de_una_factura_de_300_lineas_no_se_congela(ventana):
    from PySide6.QtWidgets import QLabel
    # Una sesión de antes de la 1.26: las 300 líneas, cada una su fila.
    procesadas = lote(lectura(0, lineas_iva=lineas(300), total=363.0))
    assert len(procesadas[0][1].facturas) == 300
    ventana._rutas_actuales = ["taco.pdf"]
    ventana._on_terminado(procesadas, *CLIENTE)
    ventana.tabla.selectRow(0)
    ventana._refrescar_ficha()
    etiquetas = ventana.ficha.findChildren(QLabel)
    assert len(etiquetas) < 120          # antes, unas 3 por línea: más de 900
    assert any("288 líneas más" in e.text() for e in etiquetas)


# =====================================================================
# n.º 32 — Tope de tokens de salida y la segunda lectura una sola vez
# =====================================================================
from facturas_excel.localizar import pedir as PEDIR_LOCALIZAR   # noqa: E402  (antes del parche de conftest)


def test_gemini_tiene_un_tope_de_salida(monkeypatch):
    from facturas_excel.extraccion import MAX_TOKENS_SALIDA
    ex = extractor(monkeypatch, {"modelo-a": json.dumps(lectura(0))}, modo="no")
    assert 2000 <= MAX_TOKENS_SALIDA <= 16384
    assert ex._config().max_output_tokens == MAX_TOKENS_SALIDA
    ex._con_esquema = False                       # el respaldo sin esquema
    assert MAX_TOKENS_SALIDA <= ex._config().max_output_tokens <= 4 * MAX_TOKENS_SALIDA


def test_senalar_en_el_documento_tambien_tiene_tope(monkeypatch):
    from types import SimpleNamespace

    from google import genai

    from facturas_excel.extraccion import MAX_TOKENS_SALIDA
    pedidas = []

    class Cliente:
        def __init__(self, **_k):
            self.models = SimpleNamespace(generate_content=self._generar)

        def _generar(self, model, contents, config):
            pedidas.append(config)
            return SimpleNamespace(text='{"datos": []}', model_version=model,
                                   usage_metadata=None)
    monkeypatch.setattr(genai, "Client", Cliente)
    PEDIR_LOCALIZAR("clave", "modelo-a", b"img", [("Total", "121,00")])
    assert pedidas[0].max_output_tokens == MAX_TOKENS_SALIDA


def test_la_segunda_lectura_no_se_guarda_dos_veces():
    d = lectura(0)
    assert combinada(d, copy.deepcopy(d))["_lectura_2"] == {}
    otra = lectura(0, total=99.0, fecha_operacion="01/03/2026")
    otra.pop("concepto_texto")
    assert combinada(d, otra)["_lectura_2"] == {
        "total": 99.0, "fecha_operacion": "01/03/2026", "concepto_texto": None}


# =====================================================================
# n.º 33 — NIF con cifras de ancho completo; tipo de documento desconocido
# =====================================================================
def test_un_nif_de_ancho_completo_se_normaliza_al_leer():
    assert procesar.normaliza_nif("１２３４５６７８Ｚ") == "12345678Z"
    pr = lote(lectura(0, emisor_nif="Ｂ１２３４５６７４"))[0][1]
    assert pr.facturas[0].nif == "B12345674"
    assert validar(pr.facturas[0]).estado == OK


@pytest.mark.parametrize("nif", ["１２３４５６７８Z", "١٢٣٤٥٦٧٨Z", "X１２３４５６７L"],
                         ids=["ancho_completo", "arabes", "nie_ancho_completo"])
def test_un_dni_con_cifras_raras_no_sale_en_verde(nif):
    from facturas_excel.validacion import REVISAR, validar_nif
    assert not validar_nif(nif)
    resultado = validar(factura_correcta(nif=nif))
    assert resultado.estado == REVISAR
    assert any("NIF" in m for m in resultado.mensajes)


def test_un_tipo_de_documento_desconocido_sale_en_ambar():
    from facturas_excel import fiscal
    from facturas_excel.extraccion import TIPOS_DOCUMENTO
    from facturas_excel.validacion import REVISAR
    assert fiscal._TIPOS_DOCUMENTO == set(TIPOS_DOCUMENTO)
    avisos = fiscal.avisos(factura_correcta(tipo_documento="abono_raro"), "gasto")
    assert [g for _t, _c, g in avisos] == [REVISAR]
    assert "abono_raro" in avisos[0][0]
    for conocido in TIPOS_DOCUMENTO:
        textos = [t for t, _c, _g in fiscal.avisos(
            factura_correcta(tipo_documento=conocido), "venta")]
        assert not any("ningún tipo conocido" in t for t in textos)


# =====================================================================
# n.º 34 — texto.reparar sin coste exponencial
# =====================================================================
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.mark.parametrize("roto,candidato", [
    ("A\x01" * 28, "A" * 55 + "XB"),                        # alternados
    ("A\x01" * 200, "A" * 399 + "XB"),
    (("A" + "\x01" * 5) * 28, "A" * 167 + "XB"),             # tandas seguidas
    ("\x02\x03" * 10 + "\x01A" * 28, "A" * 70 + "XB"),
], ids=["28_alternados", "200_alternados", "28_tandas_de_5", "mezcla"])
def test_reparar_no_tarda_con_muchas_tandas_de_invisibles(roto, candidato):
    # En otro proceso: con la expresión regular de antes no acababa nunca (y
    # sin soltar el intérprete), así la prueba falla a los 20 s, no se cuelga.
    import subprocess
    import sys
    codigo = ("import time\nfrom facturas_excel.texto import reparar\n"
              f"i = time.perf_counter()\nr = reparar({roto!r}, [{candidato!r}])\n"
              "print(r is None, time.perf_counter() - i)")
    salida = subprocess.run([sys.executable, "-c", codigo], cwd=RAIZ, timeout=20,
                            capture_output=True, text=True, check=True).stdout.split()
    assert salida[0] == "True" and float(salida[1]) < 0.5


def test_reparar_sigue_recuperando_la_letra():
    from facturas_excel.texto import reparar
    assert reparar("JOS\x01 GARC\x01A SL", ["JOSÉ GARCÍA SL"]) == "JOSÉ GARCÍA SL"
    assert reparar("JOS\x01", ["JOSE", "JOSÉ"]) == "JOSÉ"
    assert reparar("JOS\x01", ["JOSE", "JOSU"]) is None      # dos posibles: ninguno
    assert reparar("JOS\x01\x01", ["JOSXYZ"]) is None        # no caben tres letras


def test_construir_con_el_nombre_del_cliente_roto_no_tarda():
    cliente = ("A" * 55 + "XB", "12345678Z")
    inicio = time.perf_counter()
    procesar.construir(lectura(0, receptor_nombre="A\x01" * 28, receptor_nif=None,
                               emisor_nif=None), cliente[1], cliente[0])
    assert time.perf_counter() - inicio < 0.5


# =====================================================================
# Corpus del fuzzing (casos dirigidos de «lecturas absurdas», semilla
# 20261009), copiado con datos inventados. Cada caso es un taco de hojas tal
# como las entrega el Worker (doble lectura combinada, o hoja no leída) y se
# lleva por todo el camino: cliente, lote, validación, ley, documentos,
# Excel, muestras, señalar, la ventana con su ficha y la sesión.
# =====================================================================
def _hoja(d1, tipo="doble", d2=None, msg=""):
    return {"tipo": tipo, "d1": d1, "d2": d2, "msg": msg}


def _tres(tipo="doble", d2=None, mala=None, **cambios):
    """Hojas 1 y 3 normales; la 2, con los cambios."""
    mala = lectura(1, **cambios) if mala is None else mala
    return [_hoja(lectura(0)), _hoja(mala, tipo, d2), _hoja(lectura(2))]


def _todo_nan():
    return dict(total=NAN, lineas_iva=[{"base": NAN, "tipo_iva": 21.0, "cuota_iva": NAN,
                                        "pct_requiv": None, "cuota_requiv": None}])


_INICIO = lectura(5, estado_pagina_factura="inicio", lineas_iva=[], total=None)
_FINAL_NAN = lectura(5, estado_pagina_factura="final", total=NAN, emisor_nombre=None,
                     receptor_nombre=None,
                     lineas_iva=[{"base": NAN, "tipo_iva": 21, "cuota_iva": NAN}])

CORPUS = {
    "nombre_lista": _tres(emisor_nombre=["PROVEEDOR", "UNO"]),
    "nombre_dict": _tres(emisor_nombre={"nombre": "X"}),
    "nombre_10000": _tres(emisor_nombre="PROVEEDOR LARGO " * 625),
    "nombre_control": _tres(emisor_nombre="JOS\x01 GARC\x0bIA SL"),
    "nombre_surrogate": _tres(emisor_nombre="PROVEEDOR \ud800 SL"),
    "nombre_no_caracteres": _tres(emisor_nombre="PROVEEDOR ￾￿ SL"),
    "numero_5000_cifras": _tres(num_factura="9" * 5000),
    "numero_int_enorme": _tres(num_factura=10 ** 400),
    "numero_float_nan": _tres(num_factura=NAN),
    "total_nan": _tres(total=NAN),
    "total_inf_simple": _tres("simple", total=INF),
    "total_int_400_cifras_simple": _tres("simple", total=10 ** 400),
    "total_texto_1e999": _tres(total="1e999"),
    "total_lista": _tres(total=[1, 2]),
    "suplidos_nan": _tres(suplidos=NAN),
    "todo_nan_simple": _tres("simple", **_todo_nan()),
    "suma_desborda_simple": _tres("simple", total=1e308, lineas_iva=[
        {"base": 1e308, "tipo_iva": 0.0, "cuota_iva": 0.0}] * 2),
    "importe_absurdo_coherente": _tres("simple", total=1210000000000.0, lineas_iva=[
        {"base": 1e12, "tipo_iva": 21.0, "cuota_iva": 2.1e11}]),
    "lineas_texto": _tres(lineas_iva="texto"),
    "lineas_numero_simple": _tres("simple", lineas_iva=5),
    "lineas_dict": _tres(lineas_iva={"base": 10}),
    "lineas_con_none": _tres(lineas_iva=[None]),
    "lineas_300": _tres(lineas_iva=lineas(300), total=363.0),
    "lineas_3000": _tres(lineas_iva=lineas(3000), total=3630.0),
    "lineas_anidadas_30": _tres(lineas_iva=anidado(30)),
    "valor_anidado_900": _tres(concepto_texto=anidado(900)),
    "fecha_imposible": _tres(fecha="31/02/2026"),
    "fecha_lista": _tres(fecha=["01/01/2026"]),
    "nif_9000": _tres(emisor_nif="B12345674" * 1000),
    "nif_ancho_completo": _tres(emisor_nif="Ｂ１２３４５６７４"),
    "nif_digitos_arabes": _tres(emisor_nif="١٢٣٤٥٦٧٨Z"),
    "tipo_documento_inventado": _tres(tipo_documento="abono_raro"),
    "opciones_raras": _tres(confianza=5, estado_pagina_factura=5,
                            es_bien_inversion="false", mencion_iva="exenta_rara"),
    "interna_error_lista": _tres(_error=[1, 2]),
    "interna_ultima_pagina_1e9": _tres(_ultima_pagina_consolidada=10 ** 9),
    "interna_paginas_union_triple": _tres(_paginas_union_manual=[[1, 2, 3]]),
    "interna_discrepancias_raras": _tres("simple", _discrepancias=[
        {"campo": "total", "etiqueta": None}]),
    "doble_d2_vacia": _tres("par", d2={}),
    "doble_d2_tipos_cambiados": _tres("par", d2=lectura(
        1, total="abc", lineas_iva="x", emisor_nif=[1], fecha={"a": 1},
        cuenta_gasto=[628])),
    "doble_d2_total_int_enorme": _tres("par", d2=lectura(1, total=10 ** 400)),
    "doble_d2_lineas_300": _tres("par", d2=lectura(1, lineas_iva=lineas(300))),
    "doble_d1_vacia_d2_normal": _tres("par", d2=lectura(1), mala={}),
    "excepcion_mensaje_raro": [_hoja(lectura(0)), _hoja(None, "excepcion",
                                                         msg="\x00\ud800" + "E" * 10000),
                               _hoja(lectura(2))],
    "lectura_vacia_todas": [_hoja({}), _hoja({}), _hoja({})],
    "union_inicio_final_nan": [_hoja(lectura(0)), _hoja(_INICIO), _hoja(_FINAL_NAN)],
}

GRAVES = ("Error en un bloque", "Algo ha fallado", "No se ha podido crear el Excel",
          "El archivo NO coincide", "No se ha guardado el ejemplo")


def leer_como_el_worker(hojas):
    """(imagen, origen, página, lectura) de cada hoja, como las entrega el
    Worker: las dos lecturas combinadas, o la hoja en rojo si no se leyó."""
    from facturas_excel.doble_lectura import combinar
    registros = []
    for pagina, hoja in enumerate(hojas, 1):
        if hoja["tipo"] == "excepcion":
            datos = {"emisor_nombre": None, "lineas_iva": [{}],
                     "_error": hoja["msg"][:120]}
        elif hoja["tipo"] == "simple":
            datos = combinar(hoja["d1"], None, "modelo-a", "")
        else:
            segunda = hoja["d2"]
            if hoja["tipo"] == "doble":
                try:
                    segunda = copy.deepcopy(hoja["d1"])
                except RecursionError:
                    segunda = hoja["d1"]
            datos = combinar(hoja["d1"], segunda, "modelo-a", "modelo-b")
        registros.append((b"", "taco.pdf", pagina, datos))
    return registros


def _sin_importes_imposibles(facturas):
    from facturas_excel import lote
    for f in facturas:
        for campo in lote.CAMPOS_NUMERO:
            valor = getattr(f, campo)
            assert valor is None or (math.isfinite(valor) and abs(valor) <= 1e9), \
                (campo, valor)
        for campo in lote.CAMPOS_TEXTO:
            valor = getattr(f, campo)
            assert valor is None or isinstance(valor, str), (campo, valor)


def _excel_sin_rarezas(facturas, tipos, carpeta):
    from openpyxl import load_workbook

    from facturas_excel.config_columnas import leer_config
    from facturas_excel.exportar import exportar_excel, verificar_excel
    from facturas_excel.rutas import ruta_config
    for tipo, xml in (("gasto", "gastos.xml"), ("venta", "ingresos.xml")):
        suyas = [f for f, t in zip(facturas, tipos) if t == tipo]
        if not suyas:
            continue
        config = leer_config(ruta_config(xml))
        ruta = str(carpeta / f"{tipo}.xlsx")
        exportar_excel(suyas, config, ruta)
        assert verificar_excel(suyas, config, ruta) == []
        libro = load_workbook(ruta)
        celdas = [str(x).lower() for fila in libro.active.iter_rows(values_only=True)
                  for x in fila if x is not None]
        libro.close()
        assert not [c for c in celdas if c.lstrip("-") in ("nan", "inf", "infinity")]


@pytest.mark.parametrize("hojas", list(CORPUS.values()), ids=list(CORPUS))
def test_corpus_del_fuzzing(hojas, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QDialog, QMessageBox

    from facturas_excel import (control_facturas, fiscal, localizar, lote,
                                muestras_revision, validacion)
    from facturas_excel.app import VentanaPrincipal
    dialogos = []
    for nombre in ("critical", "warning", "information", "question"):
        monkeypatch.setattr(QMessageBox, nombre, staticmethod(
            lambda *a, **k: dialogos.append(str(a[1] if len(a) > 1 else ""))
            or QMessageBox.No))
    monkeypatch.setattr(QDialog, "exec", lambda self: dialogos.append(
        type(self).__name__) or 0)
    monkeypatch.setattr(sesion, "_ruta", lambda: str(tmp_path / "sesion.pkl.gz"))
    inicio = time.monotonic()

    # Lo que hace el Worker al acabar de leer
    registros = leer_como_el_worker(hojas)
    analisis = procesar.analizar_cliente([d for *_, d in registros])
    nombre, nif = ((analisis.mejor.nombre, analisis.mejor.nif) if analisis.mejor
                   else CLIENTE)
    procesadas = procesar.preparar_lote(registros, nombre, nif)
    assert len(procesadas) >= 1
    assert {p for *_, p, _ in registros} <= {
        p for _, pr in procesadas
        for p in range(pr.pagina, pr.facturas[0].ultima_pagina_origen + 1)}

    # Lo que se hace con el lote
    facturas = [f for _, pr in procesadas for f in pr.facturas]
    tipos = [pr.tipo for _, pr in procesadas for _f in pr.facturas]
    _sin_importes_imposibles(facturas)
    for f, tipo in zip(facturas, tipos):
        lote.normalizar(copy.copy(f))
        validar(f)
        fiscal.avisos(f, tipo)
        fiscal.avisos(f, tipo, True)
        for valor in (f.total_impreso, f.base_iva, f.num_factura, f.nif):
            localizar._comparable(valor)
    validacion.encontrar_duplicados(facturas)
    control_facturas.controles_documentos(facturas, tipos)
    for _img, pr in procesadas:
        procesar.a_total_factura(pr)
    if len(registros) >= 2:
        procesar.fusionar_paginas_manual(registros[:2])
    procesar.preparar_lote(registros, "OTRA EMPRESA DE PRUEBA SA", "A12345674")
    _excel_sin_rarezas(facturas, tipos, tmp_path)
    muestras_revision.guardar_lecturas(registros)

    # En la ventana: la tabla, la ficha de cada fila y la sesión
    v = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=False)
    v._rutas_actuales = ["taco.pdf"]
    v._on_terminado(procesadas, nombre, nif, registros)
    filas = v.tabla.rowCount()
    assert filas >= 1
    for r in range(min(filas, 15)):
        v.tabla.selectRow(r)
        v._refrescar_ficha()
    v._guardar_muestra_revision()
    v._guardar_sesion()
    recuperada = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=True)
    assert recuperada.tabla.rowCount() == filas
    assert not sesion.apartada()

    assert not [d for d in dialogos if d.startswith(GRAVES)], dialogos
    assert time.monotonic() - inicio < 10        # tope por caso


@pytest.mark.parametrize("motivo,leida", [("MAX_TOKENS", True), ("STOP", False)],
                         ids=["cortada_por_el_tope", "rota_sin_mas"])
def test_una_respuesta_cortada_por_el_tope_se_pide_con_mas_sitio(monkeypatch, motivo, leida):
    """Una hoja legítima larga (muchos artículos) que no cabe en el tope no se
    paga tres veces cortada para acabar en rojo: el reintento pide con más
    sitio. Si solo viene rota (sin pasar del tope), se reintenta igual que
    siempre."""
    from types import SimpleNamespace

    from facturas_excel.extraccion import (ErrorLectura, MAX_TOKENS_AMPLIADO,
                                           MAX_TOKENS_SALIDA)
    completa = json.dumps(lectura(0))
    ex = extractor(monkeypatch, {"modelo-a": completa}, modo="no")
    topes = []

    def generar(model, contents, config):
        topes.append(config.max_output_tokens)
        cabe = config.max_output_tokens >= MAX_TOKENS_AMPLIADO
        return SimpleNamespace(
            text=completa if cabe else completa[:len(completa) // 2],
            candidates=[SimpleNamespace(finish_reason=SimpleNamespace(value=motivo))],
            model_version=model, usage_metadata=None)
    ex.client.models.generate_content = generar

    if leida:
        leido = ex.extraer(b"img", "taco.pdf", 1)
        assert leido.crudo["num_factura"] == "F26/00100"
        assert topes == [MAX_TOKENS_SALIDA, MAX_TOKENS_AMPLIADO]
    else:
        with pytest.raises(ErrorLectura):
            ex._leer_con("modelo-a", b"img", 1)
        assert topes == [MAX_TOKENS_SALIDA] * 3
