"""Clientes en estimación directa: sus apuntes van a Aplifisa como «IRPF/IVA»
(1.26.1).

Lo comprobó el usuario importando filas de prueba: Aplifisa decide el tipo de
operación SOLO por la Base IRPF (vacía -> «Sólo IVA»; con la base imponible,
aunque no lleve % ni cuota -> «IRPF/IVA»). Así que, para esos clientes, cada
línea sin retención lleva su base en esa columna; las de los demás, como
siempre.
"""

import os
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QDialog

import facturas_excel.app as modulo_app
from facturas_excel import almacen, clientes, copias, ventana_aplifisa
from facturas_excel.app import VentanaPrincipal
from facturas_excel.config_columnas import leer_config
from facturas_excel.dialogo_irpf import DialogoIrpf
from facturas_excel.dialogo_orden import DialogoOrden
from facturas_excel.exportar import (
    datos_para_aplifisa, exportar_excel, totales_del_excel, verificar_excel,
)
from facturas_excel.modelo import Factura
from facturas_excel.rutas import dir_datos, ruta_config

_app = QApplication.instance() or QApplication([])

PERSONA = "12345678Z"
NIE = "X1234567L"
SOCIEDAD = "B12345674"
ANONIMA = "A12345674"
COMUNIDAD = "E12345674"      # comunidad de bienes (dígito de control válido)


def _gasto(numero, base, pct=21.0, **campos):
    cuota = round(base * pct / 100, 2) if None not in (base, pct) else None
    datos = dict(num_factura=numero, fecha="03/09/2026",
                 nombre="PROVEEDOR PRUEBA SL", nif=SOCIEDAD, concepto="622",
                 subclave="G13", base_iva=base, pct_iva=pct, cuota_iva=cuota,
                 total_impreso=(round(base + (cuota or 0), 2)
                                if base is not None else None),
                 confianza_ia="alta", verificacion="doble")
    datos.update(campos)
    return Factura(**datos)


def _con_retencion(numero, base=1000.0, pct_irpf=15.0, **campos):
    cuota_irpf = round(base * pct_irpf / 100, 2)
    return _gasto(numero, base, base_irpf=base, pct_irpf=pct_irpf,
                  cuota_irpf=cuota_irpf,
                  total_impreso=round(base * 1.21 - cuota_irpf, 2), **campos)


def _lote():
    """Un poco de todo: lo normal, una exenta, una factura con dos tipos de
    IVA, una con retención (también de dos líneas), un suplido, una línea a
    cero, un abono y una «por el total»."""
    dos_tipos = [_gasto("F-2", 100.0, documento_id="D2", lineas_factura=2),
                 _gasto("F-2", 50.0, pct=10.0, documento_id="D2",
                        lineas_factura=2)]
    retenida_dos = [
        _gasto("F-5", 100.0, documento_id="D5", lineas_factura=2,
               base_irpf=150.0, pct_irpf=15.0, cuota_irpf=22.5),
        _gasto("F-5", 50.0, pct=10.0, documento_id="D5", lineas_factura=2)]
    suplido = [_gasto("F-6", 200.0, documento_id="D6", lineas_factura=2),
               _gasto("F-6", 30.0, pct=None, documento_id="D6",
                      lineas_factura=2, es_suplido=True)]
    return [
        _gasto("F-1", 100.0),                                    # 0
        _gasto("F-EX", 80.0, pct=0.0),                           # 1 exenta
        *dos_tipos,                                              # 2, 3
        _con_retencion("F-4"),                                   # 4
        *retenida_dos,                                           # 5, 6
        *suplido,                                                # 7, 8
        _gasto("F-0", 0.0),                                      # 9
        _gasto("AB-1", -40.0),                                   # 10 abono
        _gasto("F-T", 126.2, pct=None, iva_incluido_en_base=True),  # 11 total
    ]


def _columna(ruta, letra, desde=2):
    libro = load_workbook(ruta)
    try:
        hoja = libro.active
        return [hoja[f"{letra}{fila}"].value
                for fila in range(desde, hoja.max_row + 1)]
    finally:
        libro.close()


# --------------------------------------------------------------- clientes
def test_quien_puede_estar_en_estimacion_directa():
    for nif in (PERSONA, NIE, COMUNIDAD, "J12345674", "ES12345678Z"):
        assert clientes.puede_tributar_irpf(nif), nif
    for nif in (SOCIEDAD, ANONIMA, "G12345674", "ESB12345674"):
        assert not clientes.puede_tributar_irpf(nif), nif


def test_el_regimen_de_irpf_se_recuerda_por_nif():
    assert clientes.regimen_irpf(PERSONA) == ""            # sin decir: se pregunta
    clientes.guardar_regimen_irpf("12.345.678-z", clientes.DIRECTA, "AUTONOMO")
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA
    assert clientes.regimen_irpf(" 12345678z ") == clientes.DIRECTA
    clientes.guardar_regimen_irpf(PERSONA, clientes.OTRO)
    assert clientes.regimen_irpf(PERSONA) == clientes.OTRO
    assert clientes.regimen_irpf("") == ""
    assert clientes.nombre_guardado(PERSONA) == "AUTONOMO"


def test_con_el_prefijo_es_es_el_mismo_cliente():
    """Una factura puede traer el NIF del cliente como NIF-IVA (ES12345678Z):
    es el mismo cliente, y lo que se contestó no se vuelve a preguntar."""
    clientes.guardar_regimen_irpf(PERSONA, clientes.DIRECTA)
    assert clientes.regimen_irpf("ES" + PERSONA) == clientes.DIRECTA
    clientes.guardar_regimen_irpf("ES-" + NIE, clientes.OTRO)
    assert clientes.regimen_irpf(NIE) == clientes.OTRO
    assert clientes.regimen_irpf("ES" + COMUNIDAD) == ""     # no contestado


def test_una_sociedad_nunca_esta_en_directa():
    assert clientes.regimen_irpf(SOCIEDAD) == clientes.OTRO   # no se pregunta
    clientes.guardar_regimen_irpf(SOCIEDAD, clientes.DIRECTA)
    assert clientes.regimen_irpf(SOCIEDAD) == clientes.OTRO
    assert "regimen_irpf" not in almacen.Coleccion(
        "clientes", dir_datos()).leer_todo().get(SOCIEDAD, {})  # ni se guarda
    # Aunque una ficha lo traiga (a mano, o de otra versión): sigue «otro».
    almacen.Coleccion("clientes", dir_datos()).guardar(
        ANONIMA, {"regimen_irpf": clientes.DIRECTA})
    assert clientes.regimen_irpf(ANONIMA) == clientes.OTRO


def test_la_ficha_de_antes_sin_el_campo_y_lo_demas_de_la_ficha():
    """Una ficha de la 1.26.0 (sin régimen de IRPF) se lee como «sin decir»,
    y guardar el de IRPF no toca el de IVA ni el nombre ni lo confirmado."""
    almacen.Coleccion("clientes", dir_datos()).guardar(PERSONA, {
        "nombre": "TIENDA", "regimen_recargo": clientes.TOTAL,
        "recargo_equivalencia": True, "confirmado": True})
    assert clientes.regimen_irpf(PERSONA) == ""
    clientes.guardar_regimen_irpf(PERSONA, clientes.DIRECTA)
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA
    assert clientes.regimen_recargo(PERSONA) == clientes.TOTAL
    assert clientes.nombre_confirmado(PERSONA) == "TIENDA"
    # Y al revés: cambiar el de IVA no se lleva el de IRPF.
    clientes.guardar_regimen_recargo(PERSONA, clientes.DESGLOSE)
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA


def test_viaja_con_la_copia_de_seguridad_y_el_restaurar():
    clientes.guardar_regimen_irpf(PERSONA, clientes.DIRECTA, "AUTONOMO")
    buena = copias.hacer()
    clientes.guardar_regimen_irpf(PERSONA, clientes.OTRO, "AUTONOMO")
    copias.restaurar(buena)
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA


# ---------------------------------------------------------------- exportar
@pytest.mark.parametrize("xml", ["gastos.xml", "ingresos.xml"])
def test_en_directa_cada_linea_sin_retencion_lleva_su_base_irpf(tmp_path, xml):
    config = leer_config(ruta_config(xml))
    assert config.columnas["base_irpf"] == "G"
    lote = _lote()
    ruta = str(tmp_path / "directa.xlsx")
    exportar_excel(lote, config, ruta, estimacion_directa=True)

    base_irpf = _columna(ruta, "G")
    assert base_irpf == [
        "100,00",           # lo normal: la base imponible
        "80,00",            # exenta (0 %): la base que va en el Excel
        "100,00", "50,00",  # dos tipos de IVA: cada línea, la suya
        "1000,00",          # con retención: la suya, sin tocar
        "150,00", None,     # retención de toda la factura en la 1.ª: no se
                            # suma aparte la 2.ª (contaría dos veces)
        "200,00", None,     # el suplido no se imputa
        None,               # importe 0: nada
        "-40,00",           # abono: en negativo
        "126,20",           # por el total: lo que se exporta como base
    ]
    # Sin % ni cuota (como lo probó el usuario), salvo las retenciones.
    assert _columna(ruta, "H") == [None] * 4 + ["15"] + ["15"] + [None] * 6
    assert _columna(ruta, "I") == [None] * 4 + ["150,00", "22,50"] + [None] * 6
    # Lo demás, igual que sin directa.
    otra = str(tmp_path / "otro.xlsx")
    exportar_excel(lote, config, otra)
    for letra in "ABCDEFHIJKLMNOPQ":
        assert _columna(ruta, letra) == _columna(otra, letra), letra
    # Y el doble contraste, con el mismo criterio, no ve diferencias.
    assert verificar_excel(lote, config, ruta, estimacion_directa=True) == []
    assert verificar_excel(lote, config, otra) == []


def test_una_factura_que_solo_trae_suplidos_tambien_va_irpf_iva(tmp_path):
    """La de una gestoría con solo las tasas que pagó por el cliente: al leerla
    queda una línea sin base y la del suplido. Si el suplido no lleva su Base
    IRPF, ninguna fila la lleva y el apunte entra «Sólo IVA» (como con las
    columnas de IRPF vacías). Con honorarios, el suplido sigue sin ella."""
    config = leer_config(ruta_config("gastos.xml"))
    solo = [_gasto("S-1", None, pct=None, documento_id="DS", lineas_factura=2),
            _gasto("S-1", 45.0, pct=None, documento_id="DS", lineas_factura=2,
                   es_suplido=True)]
    con_honorarios = [
        _gasto("S-2", 100.0, documento_id="DH", lineas_factura=2),
        _gasto("S-2", 45.0, pct=None, documento_id="DH", lineas_factura=2,
               es_suplido=True)]
    lote = solo + con_honorarios
    ruta = str(tmp_path / "suplidos.xlsx")
    exportar_excel(lote, config, ruta, estimacion_directa=True)
    assert _columna(ruta, "G") == [None, "45,00", "100,00", None]
    assert verificar_excel(lote, config, ruta, estimacion_directa=True) == []
    # Sin directa, como siempre.
    otra = str(tmp_path / "otro.xlsx")
    exportar_excel(lote, config, otra)
    assert _columna(otra, "G") == [None] * 4


def test_la_base_sujeta_al_0_y_la_no_sujeta_si_van_en_el_excel(tmp_path):
    """Una línea sin IVA lleva en Base IRPF la base que vaya en el Excel: si
    la configuración da columna a la base sujeta al 0 % o a la no sujeta,
    esas; si no (como en gastos.xml), no hay base que imputar."""
    from dataclasses import replace as copiar
    propia = leer_config(ruta_config("gastos.xml"))
    ampliada = copiar(propia, columnas={**propia.columnas,
                                        "base_sujeta_cero": "R",
                                        "no_sujeta": "S"})
    lote = [_gasto("F-NS", None, pct=None, no_sujeta=70.0),
            _gasto("F-SC", None, pct=None, base_sujeta_cero=30.0),
            _gasto("F-MX", 100.0, no_sujeta=20.0)]
    ruta = str(tmp_path / "ampliada.xlsx")
    exportar_excel(lote, ampliada, ruta, estimacion_directa=True)
    assert _columna(ruta, "G") == ["70,00", "30,00", "120,00"]
    assert verificar_excel(lote, ampliada, ruta, estimacion_directa=True) == []
    ruta = str(tmp_path / "propia.xlsx")
    exportar_excel(lote, propia, ruta, estimacion_directa=True)
    assert _columna(ruta, "G") == [None, None, "100,00"]


def test_lo_que_no_esta_en_directa_queda_como_antes(tmp_path):
    """Sociedades y clientes en módulos: los datos de cada fila, los de la
    línea tal cual (lo de antes), y el archivo, idéntico al de siempre."""
    lote = _lote()
    config = leer_config(ruta_config("gastos.xml"))
    assert datos_para_aplifisa(lote, config) == [f.campos_dict() for f in lote]
    assert datos_para_aplifisa(lote, config, estimacion_directa=False) == \
        [f.campos_dict() for f in lote]
    antes, ahora = str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx")
    exportar_excel(lote, config, antes)
    exportar_excel(lote, config, ahora, estimacion_directa=False)
    with zipfile.ZipFile(antes) as a, zipfile.ZipFile(ahora) as b:
        assert a.namelist() == b.namelist()
        for nombre in a.namelist():
            if nombre != "docProps/core.xml":        # lleva la hora
                assert a.read(nombre) == b.read(nombre), nombre
    assert all(valor is None for valor in _columna(ahora, "G")[:4]) and \
        _columna(ahora, "G")[4] == "1000,00"


def test_verificar_compara_la_base_irpf_que_toca(tmp_path):
    """Si el archivo no lleva la Base IRPF que le toca (o la lleva sin
    tocarle), el doble contraste lo dice: no se importaría como debe."""
    config = leer_config(ruta_config("gastos.xml"))
    lote = [_gasto("F-1", 100.0)]
    sin = str(tmp_path / "sin.xlsx")
    exportar_excel(lote, config, sin)
    fallos = verificar_excel(lote, config, sin, estimacion_directa=True)
    assert len(fallos) == 1 and "Base IRPF" in fallos[0] and "100,00" in fallos[0]
    con = str(tmp_path / "con.xlsx")
    exportar_excel(lote, config, con, estimacion_directa=True)
    assert len(verificar_excel(lote, config, con)) == 1


def test_los_totales_del_excel_no_cambian_con_la_base_irpf(tmp_path):
    """La Base IRPF no es otra base: los totales del aviso final (base, IVA,
    retención y líneas) son los mismos que sin directa."""
    config = leer_config(ruta_config("gastos.xml"))
    lote = _lote()
    a, b = str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx")
    exportar_excel(lote, config, a)
    exportar_excel(lote, config, b, estimacion_directa=True)
    assert totales_del_excel(config, a) == totales_del_excel(config, b)
    assert totales_del_excel(config, b)["lineas"] == len(lote)


def test_en_modo_numero_tambien(tmp_path):
    config = leer_config(ruta_config("gastos.xml"))
    lote = [_gasto("F-1", 100.0), _con_retencion("F-2")]
    ruta = str(tmp_path / "n.xlsx")
    exportar_excel(lote, config, ruta, modo_numeros="numero",
                   estimacion_directa=True)
    assert _columna(ruta, "G") == [100.0, 1000.0]
    assert _columna(ruta, "H") == [None, 15]


# ------------------------------------------------- la pregunta y la ventana
class _Orden:
    """La pantalla del orden, sin ventana: apunta la línea del tipo de
    operación que le llega."""
    lineas = []

    def __init__(self, parent, tipo_operacion=""):
        _Orden.lineas.append(tipo_operacion)

    def exec(self):
        return QDialog.Accepted

    def recordar(self):
        pass

    def orden(self):
        return ventana_aplifisa.ORDEN_PDF


@pytest.fixture
def exportacion(tmp_path, monkeypatch):
    """Exportación de verdad (Excel real, con la configuración de Aplifisa)
    a un Escritorio de mentira, con la pregunta del IRPF contada."""
    _Orden.lineas = []
    monkeypatch.setattr(ventana_aplifisa, "DialogoOrden", _Orden)
    monkeypatch.setattr(
        modulo_app.archivo, "ruta_excel_consolidado",
        lambda cliente, ejercicio, tipo: str(
            tmp_path / f"{'GASTOS' if tipo == 'gasto' else 'INGRESOS'}.xlsx"))
    preguntas = []
    respuesta = {"exec": QDialog.Accepted, "elegido": clientes.DIRECTA}

    def preguntar(dialogo):
        preguntas.append(dialogo)
        return respuesta["exec"]
    monkeypatch.setattr(DialogoIrpf, "exec", preguntar)
    monkeypatch.setattr(DialogoIrpf, "elegido",
                        lambda self: respuesta["elegido"]
                        if respuesta["exec"] == QDialog.Accepted else "")

    def ventana(nif, nombre="CLIENTE PRUEBA", numeros=("F-1", "F-2")):
        v = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=False)
        v._cliente_nif, v._cliente_nombre = nif, nombre
        v._bloques = [{"nombre": "b1", "cliente": nombre, "nif": nif}]
        otro = ANONIMA if nif == SOCIEDAD else SOCIEDAD   # el proveedor
        v._anadir_fila(b"", _gasto(numeros[0], 100.0, nif=otro), "gasto",
                       "622", "G13", "", "b1")
        v._anadir_fila(b"", _con_retencion(numeros[1], nif=otro), "gasto",
                       "622", "G13", "", "b1")
        venta = _gasto("V-" + numeros[0], 300.0, nif=otro, concepto="705",
                       subclave="I01")
        v._anadir_fila(b"", venta, "venta", "705", "I01", "", "b1")
        v._revalidar_todo()
        return v
    return ventana, preguntas, respuesta, tmp_path


def test_en_directa_se_pregunta_una_vez_y_gastos_e_ingresos_van_irpf_iva(
        exportacion):
    ventana, preguntas, _respuesta, carpeta = exportacion
    v = ventana(PERSONA)
    v._exportar_todo()
    assert len(preguntas) == 1
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA
    assert _Orden.lineas == [
        "Tipo de operación en Aplifisa: IRPF/IVA (estimación directa)"]
    assert _columna(carpeta / "GASTOS.xlsx", "G") == ["100,00", "1000,00"]
    assert _columna(carpeta / "INGRESOS.xlsx", "G") == ["300,00"]
    assert "IRPF/IVA (estimación directa)" in v.banda.historial[-1]

    # La siguiente vez (otro lote del mismo cliente) ya no se pregunta.
    for nombre in ("GASTOS.xlsx", "INGRESOS.xlsx"):
        os.remove(carpeta / nombre)
    ventana(PERSONA, numeros=("F-3", "F-4"))._exportar_todo()
    assert len(preguntas) == 1
    assert _columna(carpeta / "GASTOS.xlsx", "G") == ["100,00", "1000,00"]


def test_modulos_u_otro_queda_solo_iva(exportacion):
    ventana, preguntas, respuesta, carpeta = exportacion
    respuesta["elegido"] = clientes.OTRO
    ventana(COMUNIDAD)._exportar_todo()
    assert len(preguntas) == 1
    assert clientes.regimen_irpf(COMUNIDAD) == clientes.OTRO
    assert _Orden.lineas == ["Tipo de operación en Aplifisa: Sólo IVA"]
    assert _columna(carpeta / "GASTOS.xlsx", "G") == [None, "1000,00"]
    assert _columna(carpeta / "INGRESOS.xlsx", "G") == [None]


@pytest.mark.parametrize("nif", [SOCIEDAD, ANONIMA])
def test_a_una_sociedad_no_se_le_pregunta(exportacion, nif):
    ventana, preguntas, _respuesta, carpeta = exportacion
    v = ventana(nif)
    v._exportar_todo()
    assert preguntas == []
    assert _Orden.lineas == ["Tipo de operación en Aplifisa: Sólo IVA"]
    assert _columna(carpeta / "GASTOS.xlsx", "G") == [None, "1000,00"]
    assert "Sólo IVA" in v.banda.historial[-1]


def test_cancelar_la_pregunta_cancela_la_exportacion(exportacion):
    ventana, preguntas, respuesta, carpeta = exportacion
    respuesta["exec"] = QDialog.Rejected
    v = ventana(NIE)
    v._exportar_todo()
    assert len(preguntas) == 1
    assert _Orden.lineas == []                    # ni se llega a elegir orden
    assert not list(carpeta.glob("*.xlsx"))       # no se exporta a ciegas
    assert clientes.regimen_irpf(NIE) == ""       # y se volverá a preguntar
    assert not any(r.get("ya_exportada") for r in v.filas)

    respuesta["exec"] = QDialog.Accepted
    v._exportar_todo()
    assert len(preguntas) == 2
    assert _columna(carpeta / "GASTOS.xlsx", "G") == ["100,00", "1000,00"]


def test_sin_nif_del_cliente_no_se_pregunta(exportacion):
    ventana, preguntas, _respuesta, carpeta = exportacion
    ventana("")._exportar_todo()
    assert preguntas == []
    assert "Sólo IVA" in _Orden.lineas[0] and "sin el NIF" in _Orden.lineas[0]
    assert _columna(carpeta / "GASTOS.xlsx", "G") == [None, "1000,00"]


def test_al_cerrar_no_se_pregunta(exportacion):
    ventana, preguntas, _respuesta, _carpeta = exportacion
    v = ventana(PERSONA)
    v._cerrando = True
    assert v._estimacion_directa_al_exportar() is None
    assert preguntas == []
    v._cerrando = False


# ---------------------------------------------------------- Configuración
def test_configuracion_cambia_el_regimen_de_irpf(exportacion):
    ventana, preguntas, respuesta, carpeta = exportacion
    v = ventana(PERSONA)
    acciones = [a.text() for m in v.menuBar().actions() if m.menu()
                for a in m.menu().actions()]
    assert "Régimen de IRPF de este cliente…" in acciones
    assert acciones.index("Régimen de IRPF de este cliente…") == \
        acciones.index("Régimen de IVA de este cliente…") + 1

    v._elegir_regimen_irpf()
    assert clientes.regimen_irpf(PERSONA) == clientes.DIRECTA
    assert "IRPF/IVA" in v.banda.historial[-1]
    respuesta["elegido"] = clientes.OTRO
    v._elegir_regimen_irpf()
    assert clientes.regimen_irpf(PERSONA) == clientes.OTRO
    assert len(preguntas) == 2
    # Cancelar en Configuración no cambia nada.
    respuesta["exec"] = QDialog.Rejected
    v._elegir_regimen_irpf()
    assert clientes.regimen_irpf(PERSONA) == clientes.OTRO
    # Y lo dicho en Configuración manda en la exportación, sin preguntar.
    respuesta["exec"] = QDialog.Accepted
    v._exportar_todo()
    assert len(preguntas) == 3
    assert _columna(carpeta / "GASTOS.xlsx", "G") == [None, "1000,00"]


def test_configuracion_con_una_sociedad_o_sin_cliente_no_pregunta(exportacion):
    ventana, preguntas, _respuesta, _carpeta = exportacion
    v = ventana(SOCIEDAD, nombre="TALLERES PRUEBA SL")
    v._elegir_regimen_irpf()
    assert preguntas == []
    assert "no tributa por el IRPF" in v.banda.historial[-1]
    v._cliente_nif = ""
    v._elegir_regimen_irpf()
    assert preguntas == [] and "por su NIF" in v.banda.historial[-1]


def test_configuracion_propone_lo_que_ya_tenia(exportacion, monkeypatch):
    ventana, _preguntas, _respuesta, _carpeta = exportacion
    clientes.guardar_regimen_irpf(PERSONA, clientes.DIRECTA)
    propuestos = []
    original = DialogoIrpf.__init__

    def init(self, cliente, parent=None, elegido=""):
        propuestos.append(elegido)
        original(self, cliente, parent, elegido=elegido)
    monkeypatch.setattr(DialogoIrpf, "__init__", init)
    ventana(PERSONA)._elegir_regimen_irpf()
    assert propuestos == [clientes.DIRECTA]


# ---------------------------------------------------------------- diálogos
def test_el_dialogo_pregunta_y_cada_boton_es_una_respuesta():
    dialogo = DialogoIrpf("AUTÓNOMO <PRUEBA>")
    textos = [b.text() for b in dialogo.findChildren(type(dialogo.boton_otro))]
    assert "Sí, estimación directa (IRPF/IVA)" in textos
    assert "No: módulos u otro (Sólo IVA)" in textos
    assert dialogo.elegido() == ""
    cancelado = []
    dialogo.rejected.connect(lambda: cancelado.append(1))
    dialogo.boton_otro.click()
    assert not cancelado            # «No» es una respuesta, no un «Cancelar»
    assert dialogo.result() == QDialog.Accepted
    assert dialogo.elegido() == clientes.OTRO

    otro = DialogoIrpf("AUTÓNOMO", elegido=clientes.DIRECTA)
    assert otro.boton_directa.isDefault()
    otro.boton_directa.click()
    assert otro.elegido() == clientes.DIRECTA

    cancelado = DialogoIrpf("AUTÓNOMO")
    cancelado.reject()
    assert cancelado.elegido() == ""


def test_la_pantalla_de_exportar_dice_como_iran():
    linea = "Tipo de operación en Aplifisa: IRPF/IVA (estimación directa)"
    dialogo = DialogoOrden(None, tipo_operacion=linea)
    assert dialogo.lbl_tipo_operacion.text() == linea
    assert not dialogo.lbl_tipo_operacion.isHidden()
    assert DialogoOrden(None).lbl_tipo_operacion.isHidden()


def test_el_intro_no_contesta_por_el_asesor():
    """La primera vez no hay respuesta que proponer, y el Intro no puede
    guardar una para siempre: con el foco en el primer botón, Qt lo daba por
    pulsado y el cliente quedaba «en estimación directa» sin que nadie lo
    eligiera. El Intro cancela (no se exporta); desde Configuración deja lo
    que ya tenía."""
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    for actual, esperado in (("", ""), (clientes.DIRECTA, clientes.DIRECTA),
                             (clientes.OTRO, clientes.OTRO)):
        dialogo = DialogoIrpf("AUTÓNOMO", elegido=actual)
        dialogo.show()
        _app.processEvents()
        QTest.keyClick(dialogo, Qt.Key_Return)
        assert dialogo.elegido() == esperado, actual
    assert dialogo.result() == QDialog.Accepted


def test_con_la_letra_grande_las_respuestas_se_leen_enteras():
    """Con la letra más grande (escalado de Windows), los dos botones de la
    respuesta no caben en el ancho mínimo de 520: la ventana se quedaba en
    él y cortaba «Sí, estimación directa (IRPF/IVA)». Tiene que crecer."""
    from PySide6.QtWidgets import QPushButton
    normal = _app.font()
    grande = type(normal)(normal)
    grande.setPointSizeF(normal.pointSizeF() * 1.4)
    _app.setFont(grande)
    try:
        dialogo = DialogoIrpf("AUTÓNOMO")
        dialogo.show()
        _app.processEvents()
        for boton in dialogo.findChildren(QPushButton):
            assert boton.width() >= boton.sizeHint().width(), boton.text()
        dialogo.close()
    finally:
        _app.setFont(normal)
