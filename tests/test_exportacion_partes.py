import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QDialog

from facturas_excel import ventana_aplifisa
import facturas_excel.app as modulo_app
from facturas_excel.app import VentanaPrincipal
from facturas_excel.modelo import Factura

_app = QApplication.instance() or QApplication([])


def _factura(numero):
    return Factura(
        num_factura=numero, fecha="03/09/2026", nombre="PROVEEDOR SL",
        nif="B86561412", concepto="622", base_iva=100.0, pct_iva=21.0,
        cuota_iva=21.0, total_impreso=121.0, confianza_ia="alta",
    )


def test_exporta_solo_el_excel_consolidado_en_el_escritorio(
        tmp_path, monkeypatch):
    ventana = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=False)
    ventana._cliente_nombre = "CLIENTE PRUEBA"
    ventana._bloques = [
        {"nombre": "lote_parte_01", "cliente": "CLIENTE PRUEBA", "nif": "X"},
        {"nombre": "lote_parte_02", "cliente": "CLIENTE PRUEBA", "nif": "X"},
    ]
    ventana._anadir_fila(b"", _factura("F-1"), "gasto", "622", "G13", "",
                         "lote_parte_01")
    ventana._anadir_fila(b"", _factura("F-2"), "gasto", "622", "G13", "",
                         "lote_parte_02")
    ventana._revalidar_todo()

    class Orden:
        def __init__(self, parent):
            pass

        def exec(self):
            return QDialog.Accepted

        def recordar(self):
            pass

        def orden(self):
            return ventana_aplifisa.ORDEN_PDF

    escritos = []
    limpiezas = []
    monkeypatch.setattr(ventana_aplifisa, "DialogoOrden", Orden)
    monkeypatch.setattr(
        modulo_app.archivo, "ruta_excel_consolidado",
        lambda cliente, ejercicio, tipo: str(
            tmp_path / f"{'GASTOS' if tipo == 'gasto' else 'INGRESOS'}_{cliente}.xlsx"))
    monkeypatch.setattr(
        modulo_app.archivo, "eliminar_excel_temporales",
        lambda cliente, tipo: limpiezas.append((cliente, tipo)) or ["parte.xlsx"])
    monkeypatch.setattr(ventana_aplifisa, "leer_config", lambda ruta: object())
    monkeypatch.setattr(ventana_aplifisa, "exportar_excel",
                        lambda facturas, config, ruta, **_k: escritos.append(ruta))
    monkeypatch.setattr(ventana_aplifisa, "verificar_excel", lambda *args: [])
    monkeypatch.setattr(ventana_aplifisa, "totales_del_excel",
                        lambda *args: {"base_iva": 200, "cuota_iva": 42})
    monkeypatch.setattr(modulo_app.QMessageBox, "information", lambda *args: None)

    ventana._exportar_todo()

    nombres = [os.path.basename(ruta) for ruta in escritos]
    assert nombres == ["GASTOS_CLIENTE PRUEBA.xlsx"]
    assert not any("parte_" in nombre for nombre in nombres)
    assert limpiezas == [("CLIENTE PRUEBA", "gasto")]


class _Orden:
    def __init__(self, parent):
        pass

    def exec(self):
        return QDialog.Accepted

    def recordar(self):
        pass

    def orden(self):
        return ventana_aplifisa.ORDEN_PDF


def _ventana_para_exportar(tmp_path, monkeypatch, numeros=("F-1",)):
    """Exportación de verdad (Excel real) al Escritorio de mentira."""
    ventana = VentanaPrincipal(comprobar_updates=False, restaurar_sesion=False)
    ventana._cliente_nombre = "CLIENTE PRUEBA"
    ventana._bloques = [{"nombre": "b1", "cliente": "CLIENTE PRUEBA", "nif": "X"}]
    for numero in numeros:
        ventana._anadir_fila(b"", _factura(numero), "gasto", "622", "G13", "", "b1")
    ventana._revalidar_todo()
    monkeypatch.setattr(ventana_aplifisa, "DialogoOrden", _Orden)
    monkeypatch.setattr(
        modulo_app.archivo, "ruta_excel_consolidado",
        lambda cliente, ejercicio, tipo: str(
            tmp_path / f"{'GASTOS' if tipo == 'gasto' else 'INGRESOS'}_{cliente}.xlsx"))
    criticos = []
    monkeypatch.setattr(ventana_aplifisa.QMessageBox, "critical",
                        lambda *args: criticos.append(args))
    return ventana, criticos


def test_un_excel_anterior_en_el_escritorio_no_impide_exportar(tmp_path, monkeypatch):
    """Antes el nombre era siempre el mismo: con el Excel anterior todavía en
    el Escritorio (abierto o sin importar) el nuevo no salía."""
    anterior = tmp_path / "GASTOS_CLIENTE PRUEBA.xlsx"
    anterior.write_bytes(b"el de la otra vez")
    ventana, criticos = _ventana_para_exportar(tmp_path, monkeypatch, ("F-7",))

    ventana._exportar_todo()

    assert not criticos
    assert anterior.read_bytes() == b"el de la otra vez"     # no se pisa
    nuevo = tmp_path / "GASTOS_CLIENTE PRUEBA_2.xlsx"
    assert nuevo.exists()
    from openpyxl import load_workbook
    libro = load_workbook(nuevo)
    assert "F-7" in [c.value for fila in libro.active.iter_rows() for c in fila]
    libro.close()
    assert "GASTOS_CLIENTE PRUEBA_2.xlsx" in ventana.banda.historial[-1]
    assert "no se ha tocado" in ventana.banda.historial[-1]
    from facturas_excel import historial
    apunte = historial.buscar("", _factura("F-7"), "gasto", "CLIENTE PRUEBA")
    assert apunte and apunte["excel"].endswith("GASTOS_CLIENTE PRUEBA_2.xlsx")

    # Y otra exportación más tarde (otro bloque): «_3», sin tocar los otros.
    ventana._anadir_fila(b"", _factura("F-8"), "gasto", "622", "G13", "", "b1")
    ventana._revalidar_todo()
    monkeypatch.setattr(ventana_aplifisa.QMessageBox, "warning", lambda *a: None)
    ventana._decidir_ya_exportadas = lambda por_tipo: (
        por_tipo.__setitem__("gasto", [f for f in por_tipo["gasto"]
                                       if f.num_factura != "F-7"]) or True)
    ventana._exportar_todo()
    assert (tmp_path / "GASTOS_CLIENTE PRUEBA_3.xlsx").exists()
    assert nuevo.exists() and anterior.exists()


def test_si_no_se_puede_guardar_avisa_y_no_queda_nada_a_medias(tmp_path, monkeypatch):
    from facturas_excel import historial
    ventana, criticos = _ventana_para_exportar(tmp_path, monkeypatch)

    def falla(facturas, config, ruta, **_k):
        open(ruta, "wb").close()                 # empezó a escribir…
        raise PermissionError(13, "Permiso denegado", ruta)
    monkeypatch.setattr(ventana_aplifisa, "exportar_excel", falla)

    ventana._exportar_todo()

    assert len(criticos) == 1
    titulo, texto = criticos[0][1], criticos[0][2]
    assert titulo == "No se ha podido crear el Excel"
    assert "No se ha exportado nada" in texto and "Permiso denegado" in texto
    assert not list(tmp_path.glob("*.xlsx"))      # lo creado a medias, fuera
    assert not historial.buscar("", _factura("F-1"), "gasto", "CLIENTE PRUEBA")


def test_un_excel_que_no_cuadra_no_se_queda_con_el_nombre_bueno(tmp_path, monkeypatch):
    """Si el Excel releído no coincide, se aparta como «NO IMPORTAR - …»: no
    se puede importar por error ni deja al bueno con «_2»."""
    ventana, criticos = _ventana_para_exportar(tmp_path, monkeypatch)
    verificar = ventana_aplifisa.verificar_excel
    monkeypatch.setattr(ventana_aplifisa, "verificar_excel",
                        lambda *a: ["línea 2: el nº no coincide"])
    ventana._exportar_todo()
    assert criticos and "NO IMPORTAR - GASTOS_CLIENTE PRUEBA.xlsx" in criticos[-1][2]
    assert sorted(p.name for p in tmp_path.glob("*.xlsx")) == [
        "NO IMPORTAR - GASTOS_CLIENTE PRUEBA.xlsx"]

    # Corregido el problema, el bueno sale con su nombre de siempre.
    monkeypatch.setattr(ventana_aplifisa, "verificar_excel", verificar)
    criticos.clear()
    ventana._exportar_todo()
    assert not criticos
    assert (tmp_path / "GASTOS_CLIENTE PRUEBA.xlsx").exists()


def test_si_falla_archivar_el_excel_bueno_se_anuncia_igual(tmp_path, monkeypatch):
    """El Excel ya está escrito, comprobado y apuntado: un fallo al poner al
    día el archivo del cliente no puede hacer creer que no se exportó."""
    ventana, criticos = _ventana_para_exportar(tmp_path, monkeypatch)

    def rota(*_a, **_k):
        raise OSError(53, "No se encuentra la ruta de red")
    monkeypatch.setattr(ventana, "_archivar_exportacion", rota)
    ventana._exportar_todo()
    assert not criticos
    final = ventana.banda.historial[-1]
    assert "Exportación terminada y comprobada" in final
    assert "GASTOS_CLIENTE PRUEBA.xlsx" in final
    assert "no se pudo poner al día el archivo del cliente" in final


def test_si_no_se_puede_quitar_lo_creado_se_dice_cual(tmp_path, monkeypatch):
    """En Windows un archivo en uso no se deja borrar ni renombrar."""
    ruta = tmp_path / "GASTOS_CLIENTE PRUEBA.xlsx"
    ruta.write_bytes(b"a medias")

    def en_uso(*_a):
        raise PermissionError(32, "en uso por otro proceso")
    with monkeypatch.context() as m:
        m.setattr(ventana_aplifisa.os, "remove", en_uso)
        m.setattr(ventana_aplifisa.os, "replace", en_uso)
        texto = ventana_aplifisa.AplifisaMixin._retirar_excel_fallidos(
            [str(ruta)], borrar=True)
    assert "No se ha podido quitar «GASTOS_CLIENTE PRUEBA.xlsx»" in texto
    assert "NO lo importe" in texto
    # Si sí se puede, desaparece sin más.
    assert ventana_aplifisa.AplifisaMixin._retirar_excel_fallidos(
        [str(ruta)], borrar=True) == ""
    assert not ruta.exists()


def test_un_fallo_que_no_es_del_escritorio_no_culpa_al_escritorio(tmp_path, monkeypatch):
    from facturas_excel.rutas import dir_datos
    ventana, criticos = _ventana_para_exportar(tmp_path, monkeypatch)
    monkeypatch.setattr(ventana_aplifisa, "leer_config",
                        lambda ruta: (_ for _ in ()).throw(ValueError("XML de columnas roto")))
    ventana._exportar_todo()
    texto = criticos[-1][2]
    assert "XML de columnas roto" in texto and "Escritorio" not in texto
    assert "errores.log" in texto
    with open(os.path.join(dir_datos(), "errores.log"), encoding="utf-8") as fh:
        assert "XML de columnas roto" in fh.read()
